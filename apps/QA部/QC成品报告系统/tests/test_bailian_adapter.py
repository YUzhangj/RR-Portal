import base64
import copy
from dataclasses import replace
import io
import json
import re
from pathlib import Path
import struct
import tempfile
import unittest
import uuid
from unittest.mock import patch
import zipfile

import openpyxl
from PIL import Image
from reportlab.pdfgen.canvas import Canvas
from pypdf import PdfReader

import bailian_adapter as adapter
from ai_service import AIConfig, AIConfigurationError, AIInputError, AIResponseError, AIService, AIValidationError, _mock_po_data


def chat_response(value, finish="stop"):
    return {"id": "chatcmpl-fixture", "model": "qwen3.8-flash", "choices": [
        {"finish_reason": finish, "message": {"role": "assistant", "content": json.dumps(value)}}],
        "usage": {"prompt_tokens": 11, "completion_tokens": 12, "total_tokens": 23}}


class Response(io.BytesIO):
    headers = {"x-request-id": "req-bailian-fixture"}


def xls_fixture():
    """A small genuine BIFF8 workbook, constructed without an extra writer dependency."""
    record = lambda kind, body: struct.pack("<HH", kind, len(body)) + body
    bof = lambda kind: record(0x0809, struct.pack("<HHHHII", 0x0600, kind, 0x0DBB, 0x07CC, 0, 6))
    text = b"SYNTHETIC-XLS-MARKER"
    sst = record(0x00FC, struct.pack("<IIHB", 1, 1, len(text), 0) + text)
    xf = record(0x00E0, struct.pack("<HH", 0, 0) + b"\0" * 16)
    eof = record(0x000A, b"")
    name = b"PO"
    bound = lambda offset: record(0x0085, struct.pack("<IBBBB", offset, 0, 0, len(name), 0) + name)
    globals_ = bof(5) + xf + sst + bound(0) + eof
    globals_ = bof(5) + xf + sst + bound(len(globals_)) + eof
    return globals_ + bof(0x10) + record(0x0200, struct.pack("<IIHHH", 0, 1, 0, 2, 0)) + \
        record(0x00FD, struct.pack("<HHHI", 0, 0, 0, 0)) + \
        record(0x0203, struct.pack("<HHHd", 0, 1, 0, 7)) + eof


class BailianAdapterTest(unittest.TestCase):
    def service(self, **kwargs):
        return AIService(AIConfig(api_key="synthetic-key", model="qwen3.8-flash", api_format="bailian_chat",
                                  base_url="https://workspace.cn-beijing.maas.aliyuncs.com/compatible-mode/v1",
                                  max_output_tokens=64000, **kwargs))

    def extract(self, suffix, content, result=None):
        captured = []
        def post(req, timeout):
            captured.append((req, timeout))
            return Response(json.dumps(chat_response(result or _mock_po_data(source_kind="spreadsheet" if suffix in (".xlsx", ".xls", ".csv") else "page"))).encode())
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / ("source" + suffix)
            path.write_bytes(content)
            with patch("ai_service.request.urlopen", side_effect=post):
                value = self.service().extract_po(path)
        self.assertEqual(len(captured), 1)
        return value, json.loads(captured[0][0].data), captured[0][0]

    def test_explicit_protocol_default_and_validation(self):
        self.assertEqual(AIConfig().api_format, "responses")
        self.assertTrue(AIConfig().responses_url.endswith("/responses"))
        self.assertTrue(self.service().config.responses_url.endswith("/chat/completions"))
        with patch.dict("os.environ", {"OPENAI_API_FORMAT": "bailian_chat"}):
            self.assertEqual(AIConfig.from_env().api_format, "bailian_chat")
        with self.assertRaises(AIConfigurationError):
            AIConfig(api_key="test", api_format="guess").validate()

    def test_pdf_bytes_survive_native_file_and_full_schema_is_in_prompt(self):
        buffer = io.BytesIO()
        canvas = Canvas(buffer)
        marker = "SYNTHETIC-PDF-" + uuid.uuid4().hex
        canvas.drawString(50, 700, marker)
        canvas.save()
        content = buffer.getvalue()
        value, payload, req = self.extract(".pdf", content)
        file = payload["messages"][1]["content"][1]
        self.assertEqual(file["type"], "file")
        self.assertEqual(base64.b64decode(file["file"]["file_data"].split(",", 1)[1]), content)
        self.assertIn(marker, PdfReader(io.BytesIO(base64.b64decode(file["file"]["file_data"].split(",", 1)[1]))).pages[0].extract_text())
        self.assertEqual(file["file"]["filename"], "source.pdf")
        self.assertNotIn("text", payload)
        self.assertEqual(payload["response_format"], {"type": "json_object"})
        self.assertIn('"additionalProperties":false', payload["messages"][0]["content"])
        self.assertIn("document_fields", payload["messages"][0]["content"])
        self.assertFalse(payload["enable_thinking"])
        self.assertEqual(payload["max_completion_tokens"], 64000)
        self.assertEqual(req.full_url, self.service().config.responses_url)
        self.assertEqual(value.metadata.provider, "bailian")
        self.assertEqual(value.metadata.usage, {"input_tokens": 11, "output_tokens": 12, "total_tokens": 23})
        self.assertTrue(value.data["document_fields"]["factory"]["needs_review"])

    def test_png_bytes_are_preserved_and_original_detail_is_not_sent(self):
        buffer = io.BytesIO()
        Image.new("RGB", (8, 8), (33, 60, 100)).save(buffer, "PNG")
        value, payload, _ = self.extract(".png", buffer.getvalue())
        part = payload["messages"][1]["content"][1]
        self.assertEqual(part["type"], "image_url")
        self.assertEqual(base64.b64decode(part["image_url"]["url"].split(",", 1)[1]), buffer.getvalue())
        self.assertNotIn("detail", part["image_url"])

    def test_photo_workflow_uses_chat_and_preserves_existing_review_constraints(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "photo.png"
            Image.new("RGB", (8, 8), (12, 30, 80)).save(path)
            photos = [{"path": str(path), "photo_id": "7", "slot_key": "barcode"}]
            data = AIService(replace(self.service().config, mock_mode=True)).inspect_photos(photos).data
            captured = []
            def post(req, timeout):
                captured.append(json.loads(req.data))
                return Response(json.dumps(chat_response(data)).encode())
            with patch("ai_service.request.urlopen", side_effect=post):
                value = self.service().inspect_photos(photos)
            self.assertEqual(len(captured), 1)
            self.assertEqual(captured[0]["messages"][1]["content"][-1]["type"], "image_url")
            self.assertEqual(value.metadata.provider, "bailian")
            self.assertTrue(value.data["findings"][0]["requires_qc_count"])

    def test_csv_keeps_coordinates_empty_rows_and_formula_text_without_evaluation(self):
        value, payload, _ = self.extract(".csv", b'PO,00123\n\n"=1+1",7\n')
        text = payload["messages"][1]["content"][1]["text"]
        rows = [json.loads(line) for line in text.splitlines()[1:]]
        self.assertIn({"sheet": "CSV", "row": 2}, rows)
        self.assertTrue(any(r.get("cell") == "B1" and r["value"] == "00123" for r in rows))
        self.assertTrue(any(r.get("cell") == "A3" and r["value"] == "=1+1" for r in rows))
        self.assertIn("NOT evaluated", text)
        self.assertEqual(value.metadata.provider, "bailian")

    def test_real_xlsx_keeps_sheet_cell_format_and_unevaluated_formula(self):
        book = openpyxl.Workbook()
        page = book.active
        page.title = "订单PO"
        page["A1"] = "SYNTHETIC-EXCEL-MARKER"
        page["B2"] = 123
        page["B2"].number_format = "000000"
        page["C3"] = "=SUM(B2,1)"
        buffer = io.BytesIO()
        book.save(buffer)
        value, payload, _ = self.extract(".xlsx", buffer.getvalue())
        rows = [json.loads(line) for line in payload["messages"][1]["content"][1]["text"].splitlines()[1:]]
        self.assertTrue(any(r.get("sheet") == "订单PO" and r.get("cell") == "B2" and r["value"] == 123 and r["number_format"] == "000000" for r in rows))
        self.assertTrue(any(r.get("cell") == "C3" and r["value"] == "=SUM(B2,1)" and r["kind"] == "f" for r in rows))

    def test_real_biff_xls_keeps_stored_values_and_cell_addresses(self):
        text = adapter.spreadsheet_text(xls_fixture(), ".xls")
        rows = [json.loads(line) for line in text.splitlines()[1:]]
        self.assertTrue(any(r.get("cell") == "A1" and r["value"] == "SYNTHETIC-XLS-MARKER" for r in rows))
        self.assertTrue(any(r.get("cell") == "B1" and r["value"] == 7 and r["number_format"] == "General" for r in rows))
        self.assertIn('"formulas_evaluated":false', text)

    def test_xlsx_stream_reads_real_cells_regardless_of_declared_dimension(self):
        book = openpyxl.Workbook()
        book.active["A1"] = "FIRST"
        book.active["C3"] = "MUST_NOT_SKIP"
        original = io.BytesIO()
        book.save(original)
        for dimension in [b'<dimension ref="A1"/>', b'', b'<dimension ref="A1:XFD1048576"/>']:
            with self.subTest(dimension=dimension):
                modified = io.BytesIO()
                with zipfile.ZipFile(io.BytesIO(original.getvalue())) as source, zipfile.ZipFile(modified, "w") as target:
                    for info in source.infolist():
                        data = source.read(info.filename)
                        if info.filename == "xl/worksheets/sheet1.xml":
                            data, count = re.subn(rb'<dimension\b[^>]*/>', lambda _: dimension, data)
                            self.assertEqual(count, 1)
                        target.writestr(info, data)
                content = modified.getvalue()
                rows = [json.loads(line) for line in adapter.spreadsheet_text(content, ".xlsx").splitlines()[1:]]
                self.assertTrue(any(r.get("cell") == "C3" and r["value"] == "MUST_NOT_SKIP" for r in rows))
                self.assertTrue(any(r.get("cell") == "A1" and r["value"] == "FIRST" for r in rows))
                for limit, maximum in [("MAX_ROWS", 2), ("MAX_COLUMNS", 2), ("MAX_CELLS", 1)]:
                    with self.subTest(limit=limit), patch.object(adapter, limit, maximum):
                        with self.assertRaises(adapter.AdapterInputError):
                            adapter.spreadsheet_text(content, ".xlsx")

    def test_spreadsheet_limit_failures_never_return_partial_content(self):
        for limit, maximum, content in [("MAX_ROWS", 1, b"a\nb\n"), ("MAX_COLUMNS", 1, b"a,b\n"),
                                        ("MAX_CELLS", 1, b"a,b\n"), ("MAX_TEXT_BYTES", 10, b"a\n")]:
            with self.subTest(limit=limit), patch.object(adapter, limit, maximum):
                with self.assertRaises(adapter.AdapterInputError):
                    adapter.spreadsheet_text(content, ".csv")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "bad.xlsx"
            path.write_bytes(b"PK\x03\x04broken")
            with patch("ai_service.request.urlopen") as post:
                with self.assertRaises(AIInputError):
                    self.service().extract_po(path)
            post.assert_not_called()

    def test_xlsx_zip_guard_rejects_bomb_and_traversal_before_workbook_parse(self):
        for name in ["../unsafe.xml", "xl/sharedStrings.xml"]:
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr(name, "a" * 200)
            with patch.object(adapter, "MAX_ZIP_ENTRY_BYTES", 100), patch("openpyxl.load_workbook") as load:
                with self.assertRaises(adapter.AdapterInputError):
                    adapter.spreadsheet_text(buffer.getvalue(), ".xlsx")
            load.assert_not_called()

    def test_chat_rejects_length_refusal_tools_empty_or_invalid_json(self):
        cases = []
        for finish in ["length", "content_filter", "tool_calls", None]:
            cases.append(chat_response({"ok": True}, finish))
        for key, value in [("refusal", "private refusal"), ("tool_calls", [{"function": "private"}]), ("content", ""), ("content", "{broken"), ("content", "[]"), ("content", '{"x":NaN}')]:
            response = chat_response({"ok": True})
            response["choices"][0]["message"][key] = value
            cases.append(response)
        for response in cases:
            with self.subTest(response=response):
                with self.assertRaises(adapter.AdapterResponseError):
                    adapter.normalize_response(response)

    def test_schema_invalid_successful_json_still_fails_local_validation(self):
        with self.assertRaises(AIValidationError):
            self.extract(".pdf", b"%PDF-1.4\n% synthetic", {"unexpected": "not a PO result"})

    def test_actual_json_byte_budget_blocks_before_any_post(self):
        service = self.service()
        with patch("bailian_adapter.request_limit", return_value=10), patch("ai_service.request.urlopen") as post:
            with self.assertRaises(AIInputError):
                service._post_responses({"unicode": "订单" * 20})
        post.assert_not_called()
        self.assertEqual(adapter.request_limit("https://dashscope.aliyuncs.com/compatible-mode/v1"), 16 * 1024 * 1024)
        self.assertEqual(adapter.request_limit("https://w.cn-beijing.maas.aliyuncs.com/compatible-mode/v1"), 64 * 1024 * 1024)


if __name__ == "__main__":
    unittest.main()
