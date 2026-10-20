"""Synthetic regressions; writable application state is isolated before import."""
import copy
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from unittest.mock import patch
import zipfile
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
ISOLATED = tempfile.TemporaryDirectory(prefix='zuru-962-review-')
os.environ['DATA_PATH'] = ISOLATED.name
os.environ['LOCAL_MODE'] = '0'
sys.path.insert(0, str(ROOT))
import app as service
import po_parser
import openpyxl

ENTRY = {'cn_name':'测试品名', 'workshop':'测试车间', 'schedule_name':'测试排期名'}


def line(code='111', sku=None, customer_po='TEST-HK'):
    return {'simple_no':code, 'sku_spec':sku or code+'-S001', 'sku':code,
            'price':1.25, 'qty':10, 'delivery':datetime(2026, 11, 1), 'customer_po':customer_po}


def order(po='4500000001', destination='Hong Kong', customer='Fixture Customer', lines=None):
    return {'po_number':po, 'po_date':datetime(2026, 10, 20), 'destination':destination,
            'customer':customer, 'lines':lines if lines is not None else [line()]}


def item_workbook(rows):
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = '货号'
    sheet.append(['货号','中文名称','生产车间','排期名'])
    for row in rows: sheet.append(row)
    out = io.BytesIO(); book.save(out); book.close(); return out.getvalue()


class ReviewRegressions(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=ISOLATED.name)
        self.addCleanup(self.temp.cleanup)
        for key, name in [('UPLOAD_FOLDER','uploads'), ('EXPORT_FOLDER','exports'), ('DATA_FOLDER','data')]:
            directory = Path(self.temp.name, name); directory.mkdir()
            service.app.config[key] = str(directory)
        service.app.config['TESTING'] = True
        service.ITEM_MAP = {'111':dict(ENTRY), '222':dict(ENTRY)}
        service.CANCELLED_POS = set()
        self.client = service.app.test_client()

    def export(self, orders):
        result = service._generate_excel(orders)
        self.assertTrue(result['ok'])
        return result, openpyxl.load_workbook(result['filepath'], data_only=False)

    def cancellation(self, shapes=(), cells=()):
        file = Path(self.temp.name, 'po.xlsx')
        drawing = '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        for runs in shapes:
            drawing += '<xdr:twoCellAnchor><xdr:sp><xdr:txBody><a:p>' + ''.join('<a:r><a:t>'+escape(text)+'</a:t></a:r>' for text in runs) + '</a:p></xdr:txBody></xdr:sp></xdr:twoCellAnchor>'
        drawing += '</xdr:wsDr>'
        with zipfile.ZipFile(file, 'w') as archive:
            archive.writestr('xl/drawings/drawing1.xml', drawing)
            archive.writestr('xl/sharedStrings.xml', '<sst>'+''.join('<si><t>'+escape(text)+'</t></si>' for text in cells)+'</sst>')
        return po_parser.is_cancelled(file, filename='PO-4500000001.xlsx')

    def test_all_country_aliases_and_read_only_translate_endpoint_are_retained(self):
        aliases = json.loads((ROOT/'data/country_map.json').read_text())
        self.assertGreaterEqual(len(aliases), 713)
        for alias, expected in aliases.items():
            with self.subTest(alias=alias):
                if alias in ('USA', 'United States', 'US', 'U.S.', 'U.S.A.', 'United States of America', 'America'):
                    expected = '美国'
                self.assertEqual(service._translate_country(alias), expected)
        self.assertEqual(service._translate_country('Mexico'), '墨西哥-北美洲')
        self.assertEqual(service._translate_country('Argentina'), '阿根廷-南美洲')
        self.assertEqual(self.client.get('/api/translate').status_code, 400)
        data = self.client.get('/api/translate', query_string={'name':'Korea, Republic of'}).get_json()
        self.assertTrue(data['fully_translated']); self.assertEqual(data['output'], '南韩-亚洲')
        data = self.client.get('/api/translate', query_string={'name':'USA, Unknown Destination'}).get_json()
        self.assertFalse(data['fully_translated']); self.assertEqual(data['untranslated'], ['Unknown Destination'])
        self.assertEqual(list(Path(service.app.config['DATA_FOLDER']).iterdir()), [])

    def test_customer_prefix_returns_original_english_without_false_translation_warning(self):
        mapping = {'SYNTHETIC CUSTOMER':'United States', 'SYNTHETIC CUSTOMER - OUTLET':'USA'}
        with patch.object(service, 'CUSTOMER_COUNTRY', mapping):
            self.assertEqual(service._country_from_customer(' synthetic   customer - outlet - north '), 'USA')
            result, book = self.export([order(destination='', customer='SYNTHETIC CUSTOMER - NORTH')])
            try: self.assertEqual(book['排期']['K3'].value, '美国')
            finally: book.close()
            self.assertTrue(any('推断' in warning for warning in result['warnings']))
            self.assertFalse(any('未找到中文翻译' in warning for warning in result['warnings']))
        with patch.object(service, 'CUSTOMER_COUNTRY', {'SYNTHETIC CUSTOMER - A':'USA', 'SYNTHETIC CUSTOMER - B':'Germany'}):
            self.assertEqual(service._country_from_customer('SYNTHETIC CUSTOMER'), '')

    def test_official_comma_country_and_multi_country_line_suffixes(self):
        self.assertEqual(service._destination_parts('Korea, Republic of'), ['Korea, Republic of'])
        result, book = self.export([order(destination='Taiwan, Hong  Kong', lines=[line(customer_po='TEST-HK'), line('222', customer_po='TEST-TW')])])
        try: self.assertEqual([book['排期'].cell(r,11).value for r in (3,4)], ['香港-亚洲','台湾-亚洲'])
        finally: book.close()
        self.assertIsNone(service._country_for_line('TEST-US', ['Taiwan','Hong Kong']))
        self.assertEqual(service._country_for_line('TEST-HK', ['Hong Kong, China','Taiwan']), '香港-亚洲')

    def test_export_uses_one_item_snapshot_for_filtering_and_lookup_sheet(self):
        original = service._date_serial; switched = False
        def replace_map(value):
            nonlocal switched
            if not switched:
                switched = True
                with service._item_map_lock: service.ITEM_MAP = {'999':dict(ENTRY, workshop='新车间')}
            return original(value)
        with patch.object(service, '_date_serial', side_effect=replace_map):
            result, book = self.export([order(lines=[line(),line('222')])])
        try:
            self.assertEqual(result['written'], 2)
            self.assertEqual({str(row[0]) for row in book['货号'].iter_rows(min_row=2,values_only=True)}, {'111','222'})
            self.assertEqual(book['排期']['L3'].value, '=VLOOKUP(C3,货号!A:C,3,0)')
        finally: book.close()
        self.assertEqual(set(service.ITEM_MAP), {'999'})

    def test_simultaneous_exports_keep_distinct_complete_workbooks(self):
        class FixedDate(datetime):
            @classmethod
            def now(cls, tz=None): return cls(2026, 10, 20, 12, 0, 0)
        with patch.object(service, 'datetime', FixedDate):
            with ThreadPoolExecutor(max_workers=2) as pool:
                results = list(pool.map(lambda po: service._generate_excel([order(po=po)]), ['4500000001','4500000002']))
        self.assertNotEqual(results[0]['filename'], results[1]['filename'])
        for po, result in zip(['4500000001','4500000002'], results):
            book = openpyxl.load_workbook(result['filepath'])
            try: self.assertEqual(book['排期']['B3'].value, po)
            finally: book.close()

    def test_failed_cancel_list_replace_preserves_disk_and_memory(self):
        service.CANCELLED_POS = {'KEEP'}; service._save_cancelled_pos(service.CANCELLED_POS)
        path = Path(service._cancel_pos_file()); previous = path.read_bytes(); files = set(path.parent.iterdir())
        with patch.object(service.os, 'replace', side_effect=OSError('synthetic disk failure')):
            response = self.client.post('/api/cancelled_pos', json={'add':['NEW']})
        self.assertEqual(response.status_code, 500); self.assertEqual(service.CANCELLED_POS, {'KEEP'})
        self.assertEqual(path.read_bytes(), previous); self.assertEqual(set(path.parent.iterdir()), files)
        for invalid in ({'add':'NO'}, {'add':[1]}, {'add':['x'*101]}, []):
            self.assertEqual(self.client.post('/api/cancelled_pos', json=invalid).status_code, 400)

    def test_failed_item_map_replace_preserves_disk_memory_and_unrelated_items(self):
        path = Path(service.app.config['DATA_FOLDER'], 'item_map.json'); service._atomic_json(str(path), service.ITEM_MAP)
        previous = path.read_bytes(); memory = copy.deepcopy(service.ITEM_MAP)
        with patch.object(service.os, 'replace', side_effect=OSError('synthetic disk failure')):
            response = self.client.post('/api/update_items', data={'file':(io.BytesIO(item_workbook([['111','改名','新车间','新排期'], ['333','新货号','新车间','新排期']])),'synthetic.xlsx')})
        self.assertEqual(response.status_code, 500); self.assertEqual(service.ITEM_MAP, memory); self.assertEqual(path.read_bytes(), previous)
        self.assertEqual(sorted(file.name for file in path.parent.iterdir()), ['item_map.json'])

    def test_concurrent_item_updates_merge_and_duplicate_rows_keep_first(self):
        ready = threading.Barrier(2); original = service._read_items_from_table
        def simultaneous(path):
            result = original(path); ready.wait(timeout=5); return result
        def upload(code):
            with service.app.test_client() as client:
                return client.post('/api/update_items', data={'file':(io.BytesIO(item_workbook([[code,'第一行','A','P'],[code,'第二行','B','Q']])),'synthetic.xlsx')})
        with patch.object(service, '_read_items_from_table', side_effect=simultaneous):
            with ThreadPoolExecutor(max_workers=2) as pool: responses = list(pool.map(upload, ['333','444']))
        self.assertEqual([response.status_code for response in responses], [200,200])
        self.assertEqual(set(service.ITEM_MAP), {'111','222','333','444'})
        self.assertEqual(service.ITEM_MAP['333']['cn_name'], '第一行')
        stored = json.loads(Path(service.app.config['DATA_FOLDER'], 'item_map.json').read_text())
        self.assertEqual(stored, service.ITEM_MAP)
        self.assertEqual(responses[0].get_json()['dupes'], ['333'])

    def test_cancellation_requires_exact_independent_stamp_not_incidental_text(self):
        for text in ('Cancel Date', 'avoid shipping damage', 'Line 10 cancelled', '订单取消日期', '此前已取消的订单说明'):
            with self.subTest(text=text):
                self.assertFalse(self.cancellation(shapes=[(text,)], cells=[text]))
        for text in ('CANCEL', 'Cancelled!', '取 消', '订单取消', 'VOID'):
            with self.subTest(text=text): self.assertTrue(self.cancellation(shapes=[(text,)]))
        self.assertTrue(self.cancellation(shapes=[('Cancel Date',),('订单','取消')]))
        self.assertTrue(self.cancellation(cells=['订单取消！']))

    def test_real_workbook_parser_handles_wide_header_and_kit_subitem(self):
        book = openpyxl.Workbook(); sheet = book.active
        for cell,value in {'A6':'Date:', 'B6':'2026.10.20', 'A7':'PO #:', 'G7':'4500224166', 'A8':'Customer  Name:', 'G8':'Fixture Customer', 'U5':'Destination  Country:', 'V5':'Korea, Republic of'}.items(): sheet[cell]=value
        for col,value in enumerate(['Line','SKU','SKU-Spec','Name','Price USD','Qty','Delivery Date','Customer PO'],1): sheet.cell(14,col,value)
        for col,value in enumerate([1,'9574','9574-STD\nPRODUCT','Synthetic kit component',1.25,10,'2026.11.01','TEST-KR'],1): sheet.cell(15,col,value)
        sheet['A16']='Total'; path=Path(self.temp.name,'PO- 4500224166.xlsx');book.save(path);book.close()
        parsed=po_parser.parse(path)
        self.assertEqual(parsed['customer'],'Fixture Customer'); self.assertEqual(parsed['po_number'],'4500224166')
        self.assertEqual(parsed['destination'],'Korea, Republic of'); self.assertEqual(parsed['po_date'],datetime(2026,10,20))
        self.assertEqual(parsed['lines'][0]['sku_spec'],'9574-STDPRODUCT')
        service.ITEM_MAP={'9574':dict(ENTRY)};result,export=self.export([parsed])
        try: self.assertEqual(export['排期']['D3'].value,'9574'); self.assertEqual(export['排期']['K3'].value,'南韩-亚洲')
        finally: export.close()

    def test_process_keeps_cancelled_po_out_and_progress_task_available(self):
        service.CANCELLED_POS={'4500000001'}
        with patch.object(po_parser,'parse',return_value=order()), patch.object(po_parser,'is_cancelled',return_value=False):
            response=self.client.post('/api/process', data={'files':(io.BytesIO(item_workbook([])),'synthetic.xlsx')})
        self.assertEqual(response.status_code,200); data=response.get_json()
        self.assertEqual(data['written'],0); self.assertEqual(data['cancelled'],['synthetic.xlsx'])
        progress=self.client.get('/api/progress', query_string={'task_id':data['task_id']})
        self.assertIn('"done": true',progress.get_data(as_text=True))
        self.assertEqual(list(Path(service.app.config['UPLOAD_FOLDER']).iterdir()),[])


def tearDownModule():
    for handler in list(service.logging.getLogger().handlers):
        if handler is service._log_handler: service.logging.getLogger().removeHandler(handler); handler.close()
    ISOLATED.cleanup()


if __name__ == '__main__': unittest.main()
