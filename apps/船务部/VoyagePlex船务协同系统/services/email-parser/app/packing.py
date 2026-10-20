"""Shared mixed-carton normalization and export formulas for both factories."""
from copy import deepcopy
import math


def nonnegative_number(value, label):
    if value in (None, ""):
        return None
    try:
        result = float(value)
        if result < 0 or not math.isfinite(result):
            raise ValueError()
        return result
    except (TypeError, ValueError):
        raise ValueError(f"{label}须为非负数")


def normalized_packing_items(items):
    result = deepcopy(items)
    groups = {}
    for index, item in enumerate(result):
        group = str(item.get("packing_group") or "").strip()
        if group:
            # Source sheets may reuse their row numbers in different attachments.
            groups.setdefault((item.get("source_file", ""), group), []).append(index)
    for (_, group), indices in groups.items():
        if indices != list(range(indices[0], indices[-1] + 1)):
            raise ValueError(f"混装组 {group} 的货物须排列在一起")
        for key in ("pieces", "gross_weight", "volume", "pallet_count"):
            try:
                values = [float(result[i][key]) for i in indices if result[i].get(key) not in (None, "")]
            except (TypeError, ValueError):
                raise ValueError(f"混装组 {group} 的共用数值无效")
            if any(not math.isfinite(value) or value < 0 for value in values) or len(set(values)) > 1:
                raise ValueError(f"混装组 {group} 的共用件数、毛重、体积或卡板数不一致，请只在第一行填写")
            result[indices[0]][key] = values[0] if values else None
            for i in indices[1:]:
                result[i][key] = None
    return result, groups


def merge_packing_cells(ws, groups, first_row=6):
    for indices in groups.values():
        if len(indices) < 2:
            continue
        first, last = indices[0] + first_row, indices[-1] + first_row
        for col in (10, 11, 13):
            ws.merge_cells(start_row=first, end_row=last, start_column=col, end_column=col)
        for col in (4, 5, 14, 15, 16):
            values = [ws.cell(i + first_row, col).value for i in indices]
            if values[0] not in (None, "") and all(value in (None, "", values[0]) for value in values):
                ws.merge_cells(start_row=first, end_row=last, start_column=col, end_column=col)


def weight_formula(row, rule):
    expression = f"J{row}*P{row}"
    return f"=INT({expression})" if rule == "Truncate" else f"={expression}"


def apply_weight_formulas(ws, items, groups, rule, default_rule="Truncate", source_rules=True):
    if rule not in {"Template", "Truncate", "Exact"}:
        raise ValueError("模板毛净重公式无效")
    group_by_index = {index: indices for indices in groups.values() for index in indices}
    for index, item in enumerate(items):
        row = index + 6
        nonnegative_number(item.get("gross_weight_per_box"), "每箱毛重")
        nonnegative_number(item.get("net_weight_per_box"), "每箱净重")
        indices = group_by_index.get(index, [index])
        first_row = indices[0] + 6
        applied_rule = rule if rule != "Template" else (item.get("weight_rule", default_rule) if source_rules else default_rule)
        for col, unit_col, unit_key, total_key in ((11, "P", "gross_weight_per_box", "gross_weight"), (12, "Q", "net_weight_per_box", "net_weight")):
            if col == 11 and index != indices[0]:
                continue
            # Existing confirmed mixed groups may have an explicit common total.
            # New inference rules are pending; never invent a per-box group weight.
            units = {str(items[i].get(unit_key)) for i in indices if items[i].get(unit_key) not in (None, "")}
            if len(indices) > 1 and col == 11 and len(units) > 1:
                continue
            if len(indices) > 1 and item.get(unit_key) in (None, ""):
                continue
            box_ref = f"J{first_row}"
            unit_ref = f"{unit_col}{row}"
            expression = f"{box_ref}*{unit_ref}"
            formula = f"INT({expression})" if applied_rule == "Truncate" else expression
            if item.get(unit_key) in (None, "") or items[indices[0]].get("pieces") in (None, ""):
                formula = f'IF(OR({box_ref}="",{unit_ref}=""),"",{formula})'
            ws.cell(row, col, "=" + formula)
            fmt = item.get(f"{total_key}_format") if rule == "Template" and source_rules else None
            ws.cell(row, col).number_format = fmt or ("0" if applied_rule == "Truncate" else "0.00")


def total_pallets(items):
    counts = [nonnegative_number(item.get("pallet_count"), "卡板数") for item in items if item.get("pallet_count") not in (None, "")]
    if any(not count.is_integer() for count in counts):
        raise ValueError("卡板数须为非负整数")
    return int(sum(counts)) if counts else None


def mark_missing_weights(ws, task, rows, columns, total_row):
    from openpyxl.styles import PatternFill, Font
    from openpyxl.comments import Comment
    from openpyxl.cell.cell import MergedCell
    totals = set()
    for missing in task.get("exportMissingWeights") or []:
        index = missing["itemIndex"]
        if index not in rows:
            continue
        row = rows[index]
        for key in (missing["field"], missing["totalField"]):
            if key not in columns:
                continue
            cell = ws.cell(row, columns[key])
            if isinstance(cell, MergedCell):
                continue
            cell.value = "待核对"
            cell.fill = PatternFill("solid", fgColor="FFF200")
            cell.font = Font(color="C00000", bold=True)
            cell.comment = Comment("每箱重量缺失；此文件带缺项导出，不能作为完整重量合计使用。", "VoyagePlex")
        totals.add(missing["totalField"])
    for key in totals:
        if key in columns:
            ws.cell(total_row, columns[key], "待核对")
