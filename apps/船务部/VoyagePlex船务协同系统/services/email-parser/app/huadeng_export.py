"""Huadeng RR layouts, independent of the Xingxin forms."""
from datetime import date
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from .packing import normalized_packing_items, nonnegative_number, total_pallets
from .warehouse_truck import normalize_truck_type, recommend_truck


def _text(value):
    return str(value or "").strip()


def _reference(task, warehouse):
    details = task.get("exportDetails") or {}
    value = _text(details.get("sequenceNumber") or task.get("transportReference"))
    suffix = "车" if warehouse else "柜"
    # Task/database IDs are not physical truck or cabinet references.
    if not value:
        return f"{suffix}次待确认"
    return value if value.endswith(("车", "柜")) else value + suffix


def build_huadeng_workbook(task, base_builder=None):
    from .shipment_export import _display_date, _display_deadline, _number, _product_name_spec

    details = task.get("exportDetails") or {}
    template = task.get("exportTemplate") or {}
    mode = details.get("shipmentMode") or ("Warehouse" if task.get("warehouseGroups") else "Container")
    if mode not in {"Warehouse", "Container"}:
        raise ValueError("华登出货方式无效")
    warehouse = mode == "Warehouse"
    layout = template.get("layoutKey") or "Standard"
    if layout not in {"Standard", "CustomerItem", "CartonMarks"}:
        raise ValueError("华登普通表版式无效，请核对所选模板")
    if layout == "CartonMarks" and not warehouse:
        raise ValueError("箱唛尺寸特殊表仅适用于华登交仓任务")
    rule = template.get("weightRule") or "DisplayInteger"
    if rule not in {"Template", "DisplayInteger", "Exact", "Truncate"}:
        raise ValueError("华登毛净重公式无效")
    pickup = _text(details.get("pickupMode")) or "Normal"
    if pickup not in {"Normal", "CustomerPickup", "CustomerLoading"}:
        raise ValueError("华登出货方式须为正常安排、客自提或客上柜")
    items, groups = normalized_packing_items(list(task.get("items") or []))
    # Validate pallet counts, but never use the old palletWeight field for the pallet row.
    total_pallets(items)
    warehouse_purpose = task.get("exportVariant") == "warehouse"
    columns = [
        ("remark", "备注", 9), ("sequence", "序号", 6), ("customer", "洋行", 9),
        ("contract_number", "合同", 16), ("product_code", "货号", 17), ("name", "货名", 29),
        ("country", "国家", 10), ("category", "玩具类别", 10), ("quantity", "数量", 11),
        ("pieces", "件数", 10), ("gross_weight", "毛重", 11), ("net_weight", "净重", 11),
        ("volume", "体积", 11), ("customer_po", "客人PO#", 18),
        ("order_total_pieces", "每单总件数", 11), ("gross_weight_per_box", "每箱毛重", 11),
        ("net_weight_per_box", "每箱净重", 11),
    ]
    if layout == "CustomerItem":
        # Source container forms place Customer PO item No. at the right; warehouse forms after PO.
        columns.insert(14 if warehouse else len(columns), ("customer_item_number", "Customer PO item No.", 17))
    if layout == "CartonMarks":
        columns[0] = ("production_line", "拉", 9)
        columns[13:14] = [("dimension_length", "箱唛 CM", 9), ("dimension_width", "", 9), ("dimension_height", "", 9), ("customer_po", "客PO", 18)]
        columns = [("order_total_quantity", "每单总数量", width) if key == "order_total_pieces" else
                   ("net_net_weight_per_box", "每箱净净重", width) if key == "net_weight_per_box" else (key, label, width)
                   for key, label, width in columns]
    if warehouse_purpose:
        columns.append(("warehouse_location", "放货区", 18))
    col = {key: index for index, (key, _, _) in enumerate(columns, 1)}
    last_col = len(columns)
    wb = Workbook()
    ws = wb.active
    ws.title = "华登交仓表" if warehouse else "华登做柜表"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.2
    ws.page_margins.top = ws.page_margins.bottom = 0.35
    border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for index, (_, label, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
        cell = ws.cell(5, index, label)
        cell.font = Font(name="宋体", size=11)
        cell.alignment = center
        cell.border = border
    if layout == "CartonMarks":
        ws.merge_cells("N5:P5")

    for merged in ("A1:D1", "F1:L1", "F2:L2", "E3:M3"):
        ws.merge_cells(merged)
    ws["A1"] = f"TO：{_text(details.get('recipient')) or '余振军'}"
    pickup_label = {"Normal": "", "CustomerPickup": "客自提", "CustomerLoading": "客上柜"}[pickup]
    note = "，".join(value for value in (pickup_label, _text(details.get("customsAgent")), _text(details.get("pickupNote"))) if value)
    ws["F1"] = note or None
    if note:
        ws["F1"].fill = PatternFill("solid", fgColor="FFF200")
    reference = _reference(task, warehouse)
    ws["F2"] = f"RR {reference}" if warehouse else f"{_text(task.get('port')) or '装货港待确认'} RR {reference}"
    ws["F2"].font = Font(name="宋体", size=20)
    ws["F2"].alignment = center
    names = "、".join(dict.fromkeys(_text(group.get("warehouse")) for group in task.get("warehouseGroups") or [] if group.get("warehouse")))
    if warehouse:
        vehicle = normalize_truck_type(details.get("vehicleType") or "") or normalize_truck_type(task.get("containerType") or "")
        if not vehicle and details.get("truckTypeSource") != "manual":
            vehicle = recommend_truck(items)
        ws["E3"] = f"{_display_date(task.get('plannedShipDate'))} 入{_text(details.get('warehouse')) or names or '待确认仓库'}  {vehicle or '吨车待确认'} 吨车如下：各"
    else:
        ws["E3"] = f"{_display_date(task.get('plannedShipDate'))} 出  {_text(task.get('containerType')) or '柜型待确认'}  SO #{_text(task.get('soNumber'))}"
    for row in (1, 3):
        for cell in ws[row]:
            cell.alignment = center
            cell.font = Font(name="宋体", size=13)
    ws["A1"].alignment = left
    for row, height in ((1, 26), (2, 36), (3, 36), (5, 38)):
        ws.row_dimensions[row].height = height
    group_by_index = {index: indices for indices in groups.values() for index in indices}
    rows = {}
    current = 6
    for index, item in enumerate(items):
        row = current
        rows[index] = row
        indices = group_by_index.get(index, [index])
        first = indices[0]
        values = dict(item, remark="华登", sequence=index+1, customer=_text(task.get("customer")),
                      name=_product_name_spec(item), country=task.get("destinationCountry") or task.get("destination_country"),
                      order_total_pieces=item.get("order_total_pieces", item.get("pieces")),
                      customer_item_number=item.get("customer_item_number") or item.get("customer_po_item_no") or item.get("customer_item_no"))
        if layout == "CartonMarks":
            dims = re.split(r"\s*[*×xX]\s*", _text(item.get("box_dimensions")))
            for key, value in zip(("dimension_length", "dimension_width", "dimension_height"), dims):
                values[key] = value
            values["net_weight"] = item.get("net_net_weight")
        numeric_keys = {"quantity", "pieces", "gross_weight", "net_weight", "volume", "order_total_pieces", "order_total_quantity", "gross_weight_per_box", "net_weight_per_box", "net_net_weight_per_box"}
        for key, column in col.items():
            value = _number(values.get(key)) if key in numeric_keys else values.get(key)
            cell = ws.cell(row, column, value)
            cell.alignment = center
            cell.border = border
            cell.font = Font(name="宋体", size=11)
        ws.cell(row, col["volume"]).number_format = "0.000"
        for total_key, unit_key in (("gross_weight", "gross_weight_per_box"), ("net_weight", "net_net_weight_per_box" if layout == "CartonMarks" else "net_weight_per_box")):
            column = col[total_key]
            unit_col = col[unit_key]
            unit = nonnegative_number(item.get(unit_key), "每箱重量")
            if total_key == "gross_weight" and index != first:
                continue
            applied = item.get(f"{total_key}_rule", item.get("weight_rule", "Exact")) if rule == "Template" else rule
            ws.cell(row, column).number_format = (item.get(f"{total_key}_format") or "0") if rule == "Template" else "0.00" if rule == "Exact" else "0"
            units = {str(items[i].get(unit_key)) for i in indices if items[i].get(unit_key) not in (None, "")}
            # Confirmed mixed groups may supply a shared total instead of one common box weight.
            if len(indices) > 1 and (unit is None or total_key == "gross_weight" and len(units) > 1):
                continue
            box = f"{get_column_letter(col['pieces'])}{rows[first]}"
            weight = f"{get_column_letter(unit_col)}{row}"
            expression = f"{box}*{weight}"
            if applied == "Truncate":
                expression = f"INT({expression})"
            if unit is None or items[first].get("pieces") in (None, ""):
                expression = f'IF(OR({box}="",{weight}=""),"",{expression})'
            ws.cell(row, column, "=" + expression)
        current += 1
        if warehouse and index == indices[-1]:
            count = total_pallets([items[i] for i in indices])
            if count is not None:
                ws.cell(current, col["name"], f"{count}个卡板")
                ws.cell(current, col["gross_weight"], count)
                ws.cell(current, col["gross_weight"]).number_format = "0"
                ws.row_dimensions[current].height = 30
                current += 1
        ws.row_dimensions[row].height = 40
    # Merge only already-confirmed packing groups; no inference from product codes.
    for indices in groups.values():
        if len(indices) < 2:
            continue
        for key in ("pieces", "gross_weight", "volume", "contract_number", "product_code", "customer_po", "order_total_pieces", "gross_weight_per_box"):
            if key not in col:
                continue
            values = [ws.cell(rows[i], col[key]).value for i in indices]
            if key in ("pieces", "gross_weight", "volume") or values[0] not in (None, "") and all(value in (None, "", values[0]) for value in values):
                ws.merge_cells(start_row=rows[indices[0]], end_row=rows[indices[-1]], start_column=col[key], end_column=col[key])
    brand_row = current
    brand = _text(details.get("brand")) or "、".join(dict.fromkeys(_text(item.get("brand")) for item in items if item.get("brand")))
    if brand:
        ws.merge_cells(start_row=brand_row, end_row=brand_row, start_column=5, end_column=13)
        cell = ws.cell(brand_row, 5, "品牌：" + brand)
        cell.font = Font(name="宋体", size=16, color="FF0000")
        cell.alignment = left
    ws.row_dimensions[brand_row].height = 36
    total = brand_row + 1
    for row in range(6, total):
        for column in range(1, last_col+1):
            cell = ws.cell(row, column)
            cell.border = border
            cell.alignment = left if cell.value and row == brand_row else center
    ws.merge_cells(start_row=total, start_column=7, end_row=total, end_column=8)
    ws.cell(total, 7, "合计")
    for key in ("quantity", "pieces", "gross_weight", "net_weight", "volume"):
        column = col[key]
        letter = get_column_letter(column)
        cell = ws.cell(total, column, f"=SUM({letter}6:{letter}{brand_row})")
        cell.number_format = "0.000" if key == "volume" else ws.cell(6, column).number_format if items else "0"
        cell.font = Font(name="宋体", size=12, bold=True)
        cell.border = border
        cell.alignment = center
    ws.cell(total, 14, "CBM")
    ws.row_dimensions[total].height = 32

    def write(row, column, end, value):
        if end > column:
            ws.merge_cells(start_row=row, end_row=row, start_column=column, end_column=end)
        cell = ws.cell(row, column, value)
        cell.alignment = left
        cell.font = Font(name="宋体", size=12)
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, 28)

    write(total + 1, 1, 5, f"截数期：{_display_deadline(task.get('cutoffDate'))}")
    write(total + 2, 1, 4, f"制表：{_text(details.get('preparedBy'))} {date.today().isoformat()}")
    if warehouse:
        write(total + 2, 5, 10, f"收货人：{_text(details.get('consignee'))}")
        write(total + 2, 11, last_col, f"收货人代码：{_text(details.get('consigneeCode'))}")
        dimensions = list(dict.fromkeys(f"{_text(item.get('product_code'))}：{_text(item.get('box_dimensions'))}" for item in items if _text(item.get("box_dimensions"))))
        write(total + 3, 1, last_col, "    ".join(dimensions))
        end_row = total + 3
    else:
        write(total, 1, 5, f"SI：{_display_deadline(task.get('siDeadline'))}")
        filler = nonnegative_number(details.get("fillerWeight"), "填充物重量")
        container = nonnegative_number(details.get("containerWeight"), "柜重")
        for offset, label, value in ((1, "填充物重量", filler), (2, "柜重", container), (3, "整个集装箱重量", None)):
            write(total + offset, 8, 10, label)
            ws.cell(total + offset, 11, value)
            ws.cell(total + offset, 11).number_format = "0"
        if container is not None:
            ws.cell(total + 3, 11, f"=SUM(K{total}:K{total+2})")
        write(total + 2, 13, last_col, f"柜号：{_text(details.get('containerNumber'))}")
        write(total + 3, 13, last_col, f"船封：{_text(details.get('sealNumber'))}")
        write(total + 4, 1, last_col, f"贸易方式：{_text(details.get('tradeMode')) or '进料对口'}    发货人：东莞华登塑胶制品有限公司（4419943670）")
        write(total + 5, 1, 4, "拼箱：否")
        write(total + 5, 5, 11, f"收货人：{_text(details.get('consignee'))}")
        write(total + 5, 12, last_col, f"收货人代码：{_text(details.get('consigneeCode'))}")
        end_row = total + 5
    if _text(task.get("specialRequirements")):
        end_row += 1
        write(end_row, 5, last_col, _text(task.get("specialRequirements")))
        ws.row_dimensions[end_row].height = 42
    ws.print_area = f"A1:{get_column_letter(last_col)}{end_row}"
    if task.get("exportMissingWeights"):
        from .packing import mark_missing_weights
        mark_missing_weights(ws, task, rows, col, total)
    out = BytesIO()
    wb.save(out)
    suffix = ("仓务" if warehouse_purpose else "船务") + ("交仓表" if warehouse else "做柜表")
    safe = "".join(char if char.isalnum() or char in "-_" else "_" for char in reference)
    pending = "_待核对" if task.get("exportMissingWeights") else ""
    return out.getvalue(), f"华登_{_text(task.get('customer'))}_{safe}_{suffix}{pending}.xlsx"
