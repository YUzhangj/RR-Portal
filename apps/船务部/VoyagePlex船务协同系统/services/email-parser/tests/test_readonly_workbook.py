"""Value/date equivalence and bounded style-only metadata fixtures."""
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import tempfile
import tracemalloc
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900
from openpyxl.styles.stylesheet import Stylesheet
from app import readonly_workbook as m

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
TAG = '{' + NS + '}'


def rewrite(content, changes):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename in changes:
                value = changes[info.filename]
                if value is not None:
                    target.writestr(info.filename, value)
            else:
                target.writestr(info, source.read(info.filename))
        for name, value in changes.items():
            if name not in source.namelist() and value is not None:
                target.writestr(name, value)
    return output.getvalue()


def workbook(epoch=CALENDAR_WINDOWS_1900):
    wb = openpyxl.Workbook(); wb.epoch = epoch; ws = wb.active
    formats = ['General', 'yyyy-mm-dd', 'mm-dd-yy', 'h:mm:ss', '[h]:mm:ss',
               '0.000', '"days" 0', 'yyyy年mm月dd日', 'mm:ss.0', '[mm]:ss',
               '0.00E+00', 'dd/mm/yyyy;@', 'hh:mm:ss.000']
    for index, code in enumerate(formats, 1):
        cell = ws.cell(1, index, 45200.125); cell.number_format = code
    ws['A2'] = 'cached text'; ws['B2'] = True; ws['C2'] = '=1+2'
    output = BytesIO(); wb.save(output); wb.close()
    content = output.getvalue()
    with ZipFile(BytesIO(content)) as source:
        xml = ET.fromstring(source.read('xl/worksheets/sheet1.xml'))
    formula = next(cell for cell in xml.iter(TAG+'c') if cell.get('r') == 'C2')
    formula.find(TAG+'v').text = '3'
    return rewrite(content, {'xl/worksheets/sheet1.xml': ET.tostring(xml)})


class ValuesWorkbookTests(unittest.TestCase):
    def compare(self, content):
        old = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        new = m.load_values_workbook(BytesIO(content))
        try:
            self.assertEqual(new.epoch, old.epoch)
            self.assertEqual(new._date_formats, old._date_formats)
            self.assertEqual(new._timedelta_formats, old._timedelta_formats)
            self.assertEqual(new.sheetnames, old.sheetnames)
            for a, b in zip(old.worksheets, new.worksheets):
                left, right = list(a.values), list(b.values)
                self.assertEqual(right, left)
                self.assertEqual([[type(v) for v in row] for row in right], [[type(v) for v in row] for row in left])
            self.assertTrue(new.read_only); self.assertTrue(new.data_only)
        finally:
            old.close(); new.close()

    def test_builtin_custom_duration_1900_1904_and_cached_formula_values(self):
        for epoch in [CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904]:
            with self.subTest(epoch=epoch): self.compare(workbook(epoch))
        value = m.load_values_workbook(BytesIO(workbook()))
        try:
            row = next(value.active.values)
            self.assertIsInstance(row[1], datetime); self.assertIsInstance(row[4], timedelta)
            self.assertEqual(value.active['C2'].value, 3)
        finally: value.close()

    def test_shared_string_and_rich_text_values_still_use_original_reader(self):
        content = workbook()
        with ZipFile(BytesIO(content)) as source:
            types = ET.fromstring(source.read('[Content_Types].xml'))
            sheet = ET.fromstring(source.read('xl/worksheets/sheet1.xml'))
        ET.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
            PartName='/xl/sharedStrings.xml', ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml')
        cell = next(c for c in sheet.iter(TAG+'c') if c.get('r') == 'A2')
        for child in list(cell): cell.remove(child)
        cell.set('t', 's'); ET.SubElement(cell, TAG+'v').text = '0'
        strings = ('<sst xmlns="'+NS+'"><si><t>plain_x005F_</t><r><rPr><b/></rPr><t>中文</t></r>'
                   '<rPh sb="0" eb="1"><t>ignored phonetic</t></rPh></si></sst>').encode()
        self.compare(rewrite(content, {'[Content_Types].xml': ET.tostring(types),
            'xl/sharedStrings.xml': strings, 'xl/worksheets/sheet1.xml': ET.tostring(sheet)}))

    def test_missing_styles_and_empty_cell_xfs_match_standard(self):
        content = workbook(); self.compare(rewrite(content, {'xl/styles.xml': None}))
        with ZipFile(BytesIO(content)) as source: styles = ET.fromstring(source.read('xl/styles.xml'))
        xfs = styles.find(TAG+'cellXfs'); xfs.clear(); xfs.set('count', '0')
        self.compare(rewrite(content, {'xl/styles.xml': ET.tostring(styles)}))

    def test_reverse_order_duplicate_format_ids_and_prefixed_xml_match(self):
        content = workbook()
        with ZipFile(BytesIO(content)) as source: styles = ET.fromstring(source.read('xl/styles.xml'))
        formats = styles.find(TAG+'numFmts')
        ET.SubElement(formats, TAG+'numFmt', numFmtId='164', formatCode='[h]:mm:ss')
        styles.remove(formats); styles.append(formats)
        self.compare(rewrite(content, {'xl/styles.xml': ET.tostring(styles)}))

    def test_large_unused_fonts_are_not_constructed_or_retained(self):
        content = workbook()
        with ZipFile(BytesIO(content)) as source: styles = ET.fromstring(source.read('xl/styles.xml'))
        fonts = styles.find(TAG+'fonts'); fonts.clear(); fonts.set('count', '1')
        style = ET.tostring(styles).replace(b'<ns0:fonts count="1" />',
            b'<ns0:fonts count="50000">'+b'<ns0:font><ns0:name val="synthetic"/><ns0:sz val="11"/></ns0:font>'*50000+b'</ns0:fonts>')
        self.assertGreater(len(style), 1000000)
        content = rewrite(content, {'xl/styles.xml': style})
        with patch.object(Stylesheet, 'from_tree', side_effect=AssertionError('full stylesheet constructed')):
            tracemalloc.start()
            try:
                value = m.load_values_workbook(BytesIO(content))
                try:
                    self.assertEqual(len(value._fonts), 1); self.assertIsInstance(value.active['B1'].value, datetime)
                finally: value.close()
                _, peak = tracemalloc.get_traced_memory()
            finally: tracemalloc.stop()
        self.assertLess(peak, 8*1024*1024)

    def test_invalid_style_xml_ids_schema_depth_and_zip_sizes_fail_static(self):
        root = '<styleSheet xmlns="'+NS+'">'
        bad = [root+'<cellXfs><xf numFmtId="bad"/></cellXfs></styleSheet>',
               root+'<cellXfs><xf numFmtId="-1"/></cellXfs></styleSheet>',
               root+'<numFmts><numFmt numFmtId="164"/></numFmts></styleSheet>',
               root+'<cellXfs><wrong/></cellXfs></styleSheet>', root+'<cellXfs/><cellXfs/></styleSheet>',
               '<styleSheet xmlns="urn:wrong"/>', root+'<a>'*64+'</a>'*64+'</styleSheet>',
               '<!DOCTYPE styleSheet [<!ENTITY x "secret">]>'+root+'&x;</styleSheet>', root]
        for xml in bad:
            with self.subTest(length=len(xml)), self.assertRaisesRegex(ValueError, '^'+m._STYLE_ERROR+'$'):
                m.load_values_workbook(BytesIO(rewrite(workbook(), {'xl/styles.xml': xml.encode()})))
        class Huge:
            def getinfo(self, name): return SimpleNamespace(file_size=64*1024*1024+1, compress_size=1)
        with self.assertRaisesRegex(ValueError, '^'+m._STYLE_ERROR+'$'): m._read_value_styles(Huge(), SimpleNamespace())

    def test_failure_closes_archive_and_memory_error_is_not_downgraded(self):
        reader = m._ValuesExcelReader(BytesIO(workbook()), read_only=True, data_only=True)
        with patch.object(reader, 'read_strings', side_effect=MemoryError()):
            with self.assertRaises(MemoryError): reader.read()
        self.assertIsNone(reader.archive.fp)
        with patch.object(openpyxl, '__version__', 'unreviewed'):
            with self.assertRaises(RuntimeError): m.load_values_workbook(BytesIO(workbook()))


if __name__ == '__main__': unittest.main(verbosity=2)
