import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import Workbook
from app.product_workbook import parse_product_workbook

HEADER = ['备注','序号','洋行','合同','货号','货名','国家','玩具类别','数量','件数','毛重','净重','体积','客人PO#','每单总件数','每箱毛重','每箱净重']

def book_bytes(rows, header=None, title='839', header_row=5, extra=None):
 book=Workbook();sheet=book.active;sheet.title=title
 for _ in range(header_row-1):sheet.append([])
 sheet.append(header or HEADER)
 for row in rows:sheet.append(row)
 if extra:extra(book,sheet)
 stream=BytesIO();book.save(stream);return stream.getvalue()

def cargo(code='9298SLB-S002',name='狗蛋四代6个/箱',customer='ZURU'):
 return ['华登',1,customer,'C1',code,name,'美国','塑胶',60,10,999,888,1.5,'PO',10,2.65,1.09]

class HuadengProductWorkbookTests(unittest.TestCase):
 def test_ordinary_aliases_arbitrary_sheet_and_header_position(self):
  data=book_bytes([cargo()],header_row=9)
  row=parse_product_workbook(data,'华登普通.xlsx')[0]
  self.assertEqual(('839',10,'ZURU','9298SLB-S002','狗蛋四代',6),(row['sheetName'],row['rowNumber'],row['customer'],row['productCode'],row['productName'],row['quantityPerBox']))
  self.assertEqual((2.65,1.09),(row['grossWeightPerBox'],row['netWeightPerBox']))
  self.assertFalse(row['warnings']);self.assertNotIn('factoryRemark',row)
 def test_pallet_and_footer_do_not_stop_later_cargo_or_supply_weights(self):
  row=cargo();row[15]=row[16]=None
  other=cargo('92129SLB-S002','野生动物2个/箱');other[1]=2
  rows=parse_product_workbook(book_bytes([row,[None]*5+['3个卡板']+[None]*4+[3],other,[None]*7+['合计',999,100,5000,4000]]),'交仓.xlsx')
  self.assertEqual(2,len(rows));self.assertIsNone(rows[0]['grossWeightPerBox']);self.assertIsNone(rows[0]['netWeightPerBox']);self.assertTrue(rows[0]['warnings'])
  self.assertEqual('92129SLB-S002',rows[1]['productCode'])
 def test_carton_marks_net_net_and_dimensions(self):
  header=HEADER[:13]+['箱唛\nCM','','','客PO','每单总数量','每箱\n毛重','每箱\n净净重']
  detail=cargo('682TC','2寸小老虎400/箱','PWP')[:13]+[33.02,27.305,40.005,'PO',8,13.33,11.3]
  row=parse_product_workbook(book_bytes([detail],header),'箱唛.xlsx')[0]
  self.assertEqual('PWP',row['customer']);self.assertEqual(400,row['quantityPerBox']);self.assertEqual('33.02*27.305*40.005',row['boxDimensions'])
  self.assertIsNone(row['netWeightPerBox']);self.assertEqual(11.3,row['netNetWeightPerBox']);self.assertEqual(13.33,row['grossWeightPerBox'])
 def test_customer_item_insertion_does_not_shift_weight_columns(self):
  header=HEADER[:14]+['客 ITEM']+HEADER[14:];detail=cargo();detail.insert(14,'ITEM-001')
  row=parse_product_workbook(book_bytes([detail],header),'客ITEM.xlsx')[0]
  self.assertEqual(2.65,row['grossWeightPerBox']);self.assertEqual(1.09,row['netWeightPerBox'])
 def test_multiple_sheets_and_generic_brand_with_coded_dimensions(self):
  def extra(book,sheet):
   sheet.append([None]*5+['品牌： ZURU/RAINBOCORNS']);sheet.append(['9298SLB #:22.5*15*20.5cm']);sheet.append(['卡板：103*97*116.5'])
   other=book.create_sheet('客上柜');other.append(HEADER);other.append(cargo('92107-S001','摇篮8PCS/CTN'))
  rows=parse_product_workbook(book_bytes([cargo()],extra=extra),'多工作表.xlsx')
  self.assertEqual(2,len(rows));self.assertEqual('ZURU/RAINBOCORNS',rows[0]['brand']);self.assertEqual('22.5*15*20.5',rows[0]['boxDimensions']);self.assertEqual(8,rows[1]['quantityPerBox'])
 def test_ambiguous_or_mixed_specs_never_divide_totals(self):
  first=cargo('A','海绵宝宝24PCS/PDQ/CTN');second=cargo('B','粉色baby36PCS/6PDQ/CTN');second[1]=2
  def extra(book,sheet):sheet.merge_cells('J6:J7')
  rows=parse_product_workbook(book_bytes([first,second],extra=extra),'混装.xlsx')
  self.assertTrue(all(row['quantityPerBox'] is None for row in rows));self.assertTrue(all(any('跨行共用' in warning for warning in row['warnings']) for row in rows))
 def test_merged_identity_is_inherited_without_inventing_unit_weight(self):
  first=cargo('MSLD176','小魔法瓶3个/箱');second=cargo(None,'圣诞二代7个/箱');second[1]=2;second[2]=None;second[15]=None
  def extra(book,sheet):
   for col in ['C','E','J','P']:sheet.merge_cells(f'{col}6:{col}7')
  rows=parse_product_workbook(book_bytes([first,second],extra=extra),'混装表.xlsx')
  self.assertEqual(['MSLD176','MSLD176'],[row['productCode'] for row in rows]);self.assertEqual([3,7],[row['quantityPerBox'] for row in rows]);self.assertIsNone(rows[1]['grossWeightPerBox'])
 def test_brand_prefixes_and_conflicting_brands_need_review(self):
  def extra(book,sheet):sheet.append([None]*5+['品牌：ZURU/PetsAlive\n157开头的品牌：\nZURU/FUGGLER'])
  row=parse_product_workbook(book_bytes([cargo('15711A','灰豹96个/箱')],extra=extra),'品牌.xlsx')[0]
  self.assertEqual('ZURU/FUGGLER',row['brand'])
  def conflicting(book,sheet):sheet.append(['品牌：ONE']);sheet.append(['品牌：TWO'])
  row=parse_product_workbook(book_bytes([cargo()],extra=conflicting),'不明确品牌.xlsx')[0]
  self.assertEqual('',row['brand']);self.assertTrue(any('多个品牌' in warning for warning in row['warnings']))
 def test_xls_reads_numeric_sequence_and_customer_without_xlsx_loader(self):
  values=[[],[],[],[],HEADER,cargo(92107.0,'摇篮8个/箱','TOMY')]
  sheet=SimpleNamespace(name='621',nrows=len(values),merged_cells=[],row_values=lambda i:values[i])
  book=SimpleNamespace(sheets=lambda:[sheet],release_resources=lambda:None)
  with patch('xlrd.open_workbook',return_value=book):rows=parse_product_workbook(b'xls fixture','华登.xls')
  self.assertEqual('92107',rows[0]['productCode']);self.assertEqual('TOMY',rows[0]['customer']);self.assertEqual('621',rows[0]['sheetName'])
 def test_xlsm_supported_and_invalid_files_report_error(self):
  self.assertEqual(1,len(parse_product_workbook(book_bytes([cargo()]),'表.xlsm')))
  with self.assertRaisesRegex(ValueError,'有效'):parse_product_workbook(b'invalid','表.xls')

 def test_combined_code_name_header_and_sku_pallet_note(self):
  header=HEADER.copy();header[4]='货名';header[5]=None
  def merged(book,sheet):sheet.merge_cells('E5:F5')
  row=parse_product_workbook(book_bytes([cargo('157108SLD-S001','大毛绒75PCS/展示架')],header,extra=merged),'共用表头.xlsx')[0]
  self.assertEqual('157108SLD-S001',row['productCode']);self.assertEqual('大毛绒75PCS/展示架',row['productName']);self.assertTrue(row['warnings'])
  row=parse_product_workbook(book_bytes([cargo('15789SLB-S001-DISPC（133箱/卡板）')]),'卡板包装.xlsx')[0]
  self.assertEqual('15789SLB-S001-DISPC',row['productCode']);self.assertTrue(any('卡板包装' in warning for warning in row['warnings']))

 def test_blank_sku_header_requires_review_and_never_guesses_from_filename(self):
  header=HEADER.copy();header[4]=''
  row=parse_product_workbook(book_bytes([cargo('157108SLD-S001')],header),'不使用文件名中的货号.xlsx')[0]
  self.assertEqual('157108SLD-S001',row['productCode']);self.assertTrue(any('表头为空' in warning for warning in row['warnings']))
  header[3]='无关列'
  with self.assertRaisesRegex(ValueError,'表头'):parse_product_workbook(book_bytes([cargo()],header),'无依据.xlsx')

 def test_incomplete_identity_is_retained_for_human_correction(self):
  missing_name=cargo('A',None);missing_code=cargo(None,'猫蛋6个/箱');missing_code[1]=2
  rows=parse_product_workbook(book_bytes([missing_name,missing_code]),'待核对.xlsx')
  self.assertEqual(2,len(rows));self.assertEqual('A',rows[0]['productCode']);self.assertEqual('',rows[1]['productCode']);self.assertTrue(all(row['warnings'] for row in rows))
