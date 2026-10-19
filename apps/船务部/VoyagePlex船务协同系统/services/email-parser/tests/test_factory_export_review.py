from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles.stylesheet import Stylesheet

from app.attachments import parse_excel
from app.shipment_export import build_shipment_workbook


def parsed_workbook(workbook):
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "synthetic.xlsx"
        workbook.save(path)
        return parse_excel(path)


def task(layout="Standard"):
    return {"company": "Huadeng", "customer": "SYNTHETIC", "exportTemplate": {"layoutKey": layout},
            "exportDetails": {"shipmentMode": "Warehouse"}, "items": [{
                "product_code": "A", "product_name": "Synthetic", "pieces": 10, "quantity": 100,
                "gross_weight_per_box": 2.52, "net_weight_per_box": .72, "net_net_weight_per_box": .5,
                "order_total_quantity": 100, "pallet_count": 3, "customer_item_number": "ITEM-A",
                "box_dimensions": "10*20*30", "production_line": "Line A"}]}


class ExportImportReview(unittest.TestCase):
    def roundtrip(self, value):
        data, _ = build_shipment_workbook(value)
        return parsed_workbook(load_workbook(BytesIO(data)))

    def test_amazon_customer_item_comes_from_each_cargo_row(self):
        wb = Workbook(); ws = wb.active
        ws.append(["Spec#", "Expected Quantity", "Total Carton", "New Booking", "New Booking Key#", "Customer PO item No."])
        ws.append(["A", 100, 10, "40HQ-1", "BOOK-A", "ITEM-A"])
        ws.append(["B", 200, 20, "40HQ-1", "BOOK-A", "ITEM-B"])
        self.assertEqual(["ITEM-A", "ITEM-B"], [i["customer_item_number"] for i in parsed_workbook(wb)["items"]])

    def test_short_streamed_rows_fill_merged_trailing_identity(self):
        wb = Workbook(); ws = wb.active
        ws.append(["SKU", "Description", "Quantity", "Cartons", "Customer PO"])
        ws.append(["A", "First", 10, 1, "PO-A"])
        ws.append(["B", "Second", 20, 2])
        ws.merge_cells("E2:E3")
        self.assertEqual(["PO-A", "PO-A"], [i["customer_po"] for i in parsed_workbook(wb)["items"]])

    def test_exported_pallet_row_is_recognized_without_becoming_cargo(self):
        parsed = self.roundtrip(task())
        self.assertEqual(1, len(parsed["items"]))
        self.assertEqual(3, parsed["items"][0].get("pallet_count"))

    def test_customer_item_column_does_not_shift_unit_weights(self):
        item = self.roundtrip(task("CustomerItem"))["items"][0]
        self.assertEqual("ITEM-A", item.get("customer_item_number"))
        self.assertEqual(2.52, item["gross_weight_per_box"])
        self.assertEqual(.72, item["net_weight_per_box"])

    def test_carton_marks_preserves_its_distinct_pure_weight_and_dimensions(self):
        item = self.roundtrip(task("CartonMarks"))["items"][0]
        self.assertEqual(.5, item.get("net_net_weight_per_box"))
        self.assertEqual(2.52, item.get("gross_weight_per_box"))
        self.assertEqual("10*20*30", item.get("box_dimensions"))
        self.assertEqual(100, item.get("order_total_quantity"))
        self.assertEqual("Line A", item.get("production_line"))

    def test_mixed_explicit_totals_obey_selected_precision_without_changing_values(self):
        value = task(); first = value["items"][0]
        first.update(packing_group="mix", gross_weight=25.27, net_weight=7.21)
        first.pop("gross_weight_per_box"); first.pop("net_weight_per_box")
        value["items"].append(dict(first, product_code="B", net_weight=3.17))
        for rule, expected in (("Exact", "0.00"), ("Template", "0.000")):
            with self.subTest(rule=rule):
                current = deepcopy(value); current["exportTemplate"]["weightRule"] = rule
                for item in current["items"]:
                    item.update(gross_weight_format="0.000", net_weight_format="0.000")
                ws = load_workbook(BytesIO(build_shipment_workbook(current)[0])).active
                self.assertEqual(25.27, ws["K6"].value)
                self.assertEqual(3.17, ws["L7"].value)
                self.assertEqual(expected, ws["K6"].number_format)
                self.assertEqual(expected, ws["L7"].number_format)

    def test_source_formula_rules_and_number_formats_do_not_build_styles(self):
        wb = Workbook(); ws = wb.active
        ws.append(["备注","序号","洋行","合同","货号","货名","国家","玩具类别","数量","件数","毛重","净重","体积","客人PO#","每单总件数","每箱毛重","每箱净重"])
        ws.append(["华登",1,"TEST","C-A","A","Synthetic","US","Plastic",100,10,"=INT(J2*P2)","=J2*Q2",1,"PO-A",10,2.52,.72])
        ws["K2"].number_format = "0"; ws["L2"].number_format = "0.000"
        with patch.object(Stylesheet, "from_tree", side_effect=AssertionError("full stylesheet constructed")):
            item = parsed_workbook(wb)["items"][0]
        self.assertEqual("Truncate", item.get("weight_rule"))
        self.assertEqual("Truncate", item.get("gross_weight_rule"))
        self.assertEqual("Exact", item.get("net_weight_rule"))
        self.assertEqual("0", item.get("gross_weight_format"))
        self.assertEqual("0.000", item.get("net_weight_format"))
        value = task(); value["items"] = [item]; value["exportTemplate"]["weightRule"] = "Template"
        exported = load_workbook(BytesIO(build_shipment_workbook(value)[0])).active
        self.assertEqual("=INT(J6*P6)", exported["K6"].value)
        self.assertEqual("=J6*Q6", exported["L6"].value)
        self.assertEqual("0.000", exported["L6"].number_format)


if __name__ == "__main__":
    unittest.main()
