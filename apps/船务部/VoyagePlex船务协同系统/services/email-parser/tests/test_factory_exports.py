from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import asyncio

from openpyxl import Workbook, load_workbook
from app.main import export_shipment
from app.attachments import parse_excel
from app.huadeng_forms import parse_huadeng_form
from app.rules import parse_body
from app.shipment_export import build_shipment_workbook


class FactoryExportTests(unittest.TestCase):
    def test_xingxin_title_uses_actual_transport_reference(self):
        for mode,reference,expected in (("Container","6859柜","蛇口 ZURU 6859柜"),("Warehouse","1171","蛇口 ZURU 1171车")):
            task={"company":"Xingxin","customer":"ZURU","port":"蛇口","transportReference":reference,"exportDetails":{"shipmentMode":mode},"items":[]}
            self.assertEqual(expected,self.sheet(task)["F2"].value)

    def test_empty_header_note_has_no_yellow_fill(self):
        sheet=self.sheet({"company":"Xingxin","items":[]})
        self.assertIsNone(sheet["F1"].value)
        self.assertIsNone(sheet["F1"].fill.patternType)

    def test_xingxin_notes_use_top_strip_and_compact_bottom_right(self):
        task={"company":"Xingxin","customer":"ZURU","items":[{"supplier":"兴信"}],"specialRequirements":"兴信拖柜，丹尼送兴信拼柜，深圳报关，拉网、拍照、立放"}
        sheet=self.sheet(task)
        self.assertEqual("兴信拖柜，丹尼送兴信拼柜，深圳报关",sheet["F1"].value)
        self.assertEqual("solid",sheet["F1"].fill.patternType)
        note=next(c for row in sheet for c in row if c.value=="拉网、拍照、立放")
        self.assertEqual(14,note.column)
        self.assertTrue(any(m.min_col==14 and m.max_col==17 and note.coordinate in m for m in sheet.merged_cells.ranges))
        self.assertEqual("solid",note.fill.patternType)

    def test_long_loading_note_uses_yellow_strip_with_wrapping(self):
        text="请留意按PO/ITEM进行分批次装货，唛头标签请统一朝外，如果混装仓库会额外收分货费。"
        sheet=self.sheet({"company":"Xingxin","items":[{"supplier":"兴信"}],"specialRequirements":text})
        note=next(c for row in sheet for c in row if c.value==text.rstrip("。"))
        self.assertEqual(6,note.column)
        self.assertTrue(note.alignment.wrap_text)
        self.assertEqual("solid",note.fill.patternType)

    def test_supplier_list_stays_in_header(self):
        text="兴信拖柜，丹尼，库有，南新送兴信拼柜，深圳报关，拉网、拍照、立放"
        sheet=self.sheet({"company":"Xingxin","items":[{"supplier":"兴信"}],"specialRequirements":text})
        self.assertEqual("兴信拖柜，丹尼，库有，南新送兴信拼柜，深圳报关",sheet["F1"].value)
        self.assertIn("拉网、拍照、立放",[c.value for row in sheet for c in row])

    def test_xingxin_confirmed_weights_are_both_excel_formulas(self):
        task = {"company":"Xingxin", "customer":"ZURU", "items":[
            {"pieces":84,"gross_weight_per_box":2.52,"net_weight_per_box":.72,"gross_weight":2.52,"net_weight":.72},
            {"pieces":4,"gross_weight_per_box":4.62,"net_weight_per_box":1.7}]}
        sheet=self.sheet(task)
        for cell, expected in (("K6","=INT(J6*P6)"),("L6","=INT(J6*Q6)"),("K7","=INT(J7*P7)"),("L7","=INT(J7*Q7)")):
            self.assertEqual(expected,sheet[cell].value)
        self.assertEqual(2.52,sheet["P6"].value)
        self.assertEqual(.72,sheet["Q6"].value)
        self.assertEqual("=SUM(L6:L7)",sheet["L9"].value)

    def test_missing_unit_weights_never_export_legacy_values_as_totals(self):
        task={"company":"Xingxin","items":[{"pieces":84,"gross_weight":2.52,"net_weight":.72}]}
        sheet=self.sheet(task)
        self.assertEqual('=IF(OR(J6="",P6=""),"",INT(J6*P6))',sheet["K6"].value)
        self.assertEqual('=IF(OR(J6="",Q6=""),"",INT(J6*Q6))',sheet["L6"].value)

    def test_xingxin_template_ignores_huadeng_source_weight_rule(self):
        task={"company":"Xingxin","items":[{"pieces":84,"gross_weight_per_box":2.52,"net_weight_per_box":.72,"weight_rule":"Exact"}]}
        self.assertEqual("=INT(J6*P6)",self.sheet(task)["K6"].value)
        self.assertEqual("=INT(J6*Q6)",self.sheet(task)["L6"].value)

    def test_invalid_unit_weights_are_rejected(self):
        for key in ("gross_weight_per_box","net_weight_per_box"):
            for value in (-1, "abc", float("inf")):
                task=self.task();task["items"][0][key]=value
                with self.assertRaises(ValueError): build_shipment_workbook(task)

    def test_huadeng_does_not_call_xingxin_template(self):
        from app.huadeng_export import build_huadeng_workbook
        def forbidden(_): raise AssertionError("Huadeng must not reuse the Xingxin template")
        data,_=build_huadeng_workbook(self.task(), forbidden)
        self.assertEqual("华登交仓表",load_workbook(BytesIO(data)).active.title)

    def test_templates_reject_other_factory_even_for_same_customer(self):
        for company,other in (("Xingxin","Huadeng"),("Huadeng","Xingxin")):
            task=self.task(company);task["customer"]="ZURU";task["exportTemplate"]={"company":other}
            with self.assertRaises(ValueError): build_shipment_workbook(task)

    def test_warehouse_and_container_are_distinct_for_both_factories(self):
        for company in ("Xingxin","Huadeng"):
            task=self.task(company)
            warehouse=self.sheet(task)
            self.assertFalse(any("柜重" in str(c.value) or "船封" in str(c.value) for row in warehouse for c in row))
            task["exportDetails"]["shipmentMode"]="Container"
            task["warehouseGroups"]=[{"warehouse":"仓库"}]
            container=self.sheet(task)
            self.assertTrue(any("柜重" in str(c.value) for row in container for c in row))
    def task(self, company="Huadeng"):
        return {"company": company, "customer": "ACME Toys", "port": "盐田", "items": [
            {"product_code": "A", "pieces": 665, "gross_weight_per_box": .83,
             "gross_weight": 551.95, "quantity": 3990, "net_weight": 172.9, "pallet_count": 5}],
            "exportDetails": {"shipmentMode": "Warehouse", "consignee": "ACTUAL IMPORTER LTD.", "consigneeCode": "ABC-123", "warehouse": "盐港二号仓"}}

    def sheet(self, task):
        data, _ = build_shipment_workbook(task)
        return load_workbook(BytesIO(data)).active

    def test_actual_export_endpoint_routes_factory_and_pallets(self):
        response = asyncio.run(export_shipment(self.task()))
        self.assertEqual(200, response.status_code)
        sheet = load_workbook(BytesIO(response.body)).active
        self.assertEqual("TO：余振军", sheet["A1"].value)
        self.assertEqual("RR 车次待确认", sheet["F2"].value)
        self.assertEqual("ACME Toys", sheet["C6"].value)
        self.assertEqual("5个卡板", sheet["F7"].value)
        self.assertEqual(5,sheet["K7"].value)
        self.assertEqual("=J6*P6", sheet["K6"].value)
        self.assertEqual("=SUM(K6:K8)", sheet["K9"].value)
        self.assertTrue(any("ACTUAL IMPORTER LTD." in str(c.value) for row in sheet for c in row))

    def test_template_weight_rule_overrides_source(self):
        for company in ("Xingxin", "Huadeng"):
            task = self.task(company)
            task["exportTemplate"] = {"weightRule": "Truncate"}
            self.assertEqual("=INT(J6*P6)", self.sheet(task)["K6"].value)
            task["exportTemplate"]["weightRule"] = "Exact"
            self.assertEqual("=J6*P6", self.sheet(task)["K6"].value)

    def test_huadeng_pallet_row_uses_count_and_ignores_old_pallet_weight(self):
        task = self.task(); task["exportDetails"]["palletWeight"] = 999
        sheet = self.sheet(task)
        self.assertEqual(5, sheet["K7"].value)
        self.assertEqual("=SUM(K6:K8)", sheet["K9"].value)

    def test_template_display_precision_does_not_change_calculation(self):
        task = self.task()
        task["items"][0].update(gross_weight_format="0", net_weight_format="0")
        sheet = self.sheet(task)
        self.assertEqual("=J6*P6", sheet["K6"].value)
        self.assertEqual("0", sheet["K6"].number_format)
        self.assertEqual("0", sheet["K9"].number_format)
        self.assertEqual("0", sheet["L9"].number_format)

    def test_missing_consignee_never_falls_back_to_customer(self):
        for company in ("Xingxin", "Huadeng"):
            task = self.task(company); task["customer"] = "ZURU"
            task["exportDetails"].pop("consignee"); task["exportDetails"].pop("consigneeCode")
            sheet = self.sheet(task)
            self.assertFalse(any("ZURU INC." in str(c.value) or "66934704" in str(c.value) for row in sheet for c in row))

    def test_both_factories_merge_variants_without_double_counting(self):
        for company in ("Xingxin", "Huadeng"):
            task = self.task(company)
            first = dict(task["items"][0], packing_group="mix", product_name="第一款")
            second = dict(first, product_name="第二款", quantity=100, net_weight=20)
            task["items"] = [first, second]
            sheet = self.sheet(task)
            for merged in ("E6:E7", "J6:J7", "K6:K7", "M6:M7"):
                self.assertIn(merged, str(sheet.merged_cells))
            self.assertEqual(665, sheet["J6"].value)
            self.assertIsNone(sheet["J7"].value)
            self.assertEqual(20, sheet["L7"].value)
            self.assertEqual("5个卡板" if company=="Huadeng" else "5卡板", sheet["F8"].value)

    def test_pallet_count_validation(self):
        for value in (-1, 2.5, "abc", float("inf")):
            task = self.task(); task["items"][0]["pallet_count"] = value
            with self.assertRaises(ValueError): build_shipment_workbook(task)

    def test_mixed_variants_with_different_unit_weights_keep_group_weight(self):
        task = self.task("Xingxin")
        task["items"] = [dict(task["items"][0], packing_group="mix"),
                         dict(task["items"][0], packing_group="mix", gross_weight_per_box=.9)]
        self.assertEqual(551.95, self.sheet(task)["K6"].value)

    def test_mail_consignee_ignores_email_to_and_quoted_history(self):
        fields = parse_body("To: shipping@example.com\n收货人：ROYAL REGENT TOYS PRODUCTS LTD. 收货人代码：12345\nFrom: old@example.com\nConsignee: OLD COMPANY")
        self.assertEqual("ROYAL REGENT TOYS PRODUCTS LTD.", fields["consignee"])
        self.assertEqual("12345", fields["consignee_code"])
        self.assertEqual("", parse_body("To: shipping@example.com\n客户：ZURU")["consignee"])
        self.assertEqual("ACTUAL LTD.", parse_body("Consignee Name:\nACTUAL LTD.")["consignee"])

    def test_form_reads_pallet_row_and_actual_customer(self):
        rows = [["TO：余振军"], ["RR 510车"],
                ["备注", "序号", "洋行", "合同", "货号", "货名", "国家", "玩具类别", "数量", "件数", "毛重", "净重", "体积", "客人PO#", "每单总件数", "每箱毛重", "每箱净重"],
                ["华登", 1, "ACME", "PO", "A", "第一款", "美国", "塑胶", 3990, 665, 551.95, 172.9, 5.82, "PO", 665, .83, .26],
                [None, None, None, None, None, "5卡板", None, None, None, None, 5],
                ["收货人：ACTUAL LTD."]]
        parsed = parse_huadeng_form(rows, [])
        self.assertEqual(1, len(parsed["items"]))
        self.assertEqual(5, parsed["items"][0]["pallet_count"])
        self.assertEqual("ACME", parsed["fields"]["customer"])
        self.assertEqual("Warehouse", parsed["fields"]["shipment_mode"])
        self.assertEqual("ACTUAL LTD.", parsed["fields"]["consignee"])

    def test_generic_packing_list_reads_merged_same_sku_variants(self):
        wb = Workbook(); ws = wb.active
        ws.append(["SKU", "Description", "Quantity", "Cartons", "Case Pack", "Gross Weight", "Net Weight", "CBM"])
        ws.append(["A", "Variant one", 12, 3, 12, 10, 4, 2])
        ws.append([None, "Variant two", 24, None, 12, None, 6, None])
        for column in ("A", "D", "F", "H"): ws.merge_cells(f"{column}2:{column}3")
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "packing.xlsx"; wb.save(path)
            parsed = parse_excel(path)
        self.assertEqual(["A", "A"], [item["product_code"] for item in parsed["items"]])
        self.assertEqual(2, len(parsed["items"]))
        self.assertEqual(parsed["items"][0]["packing_group"], parsed["items"][1]["packing_group"])
        self.assertIsNone(parsed["items"][1]["pieces"])
