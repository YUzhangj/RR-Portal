import sys
import asyncio
import tempfile
import unittest
import base64
from datetime import datetime
from io import BytesIO
from email.message import EmailMessage
from pathlib import Path

import openpyxl
from docx import Document
from starlette.datastructures import UploadFile

sys.path.insert(0, str(Path(__file__).parents[1]))

from app.attachments import _clean_address, _column_map, parse_attachment
from app.eml import parse_eml
from app.main import _inventory_paths, _match_inventory_paths, _scan_inventory_paths, build_amazon_shipment_groups, build_warehouse_groups, parse_email_batch
from app.rules import filter_items_for_email, parse_body
from app.spreadsheets import parse_business_spreadsheet
from app.shipment_export import build_completed_shipment_summary_workbook, build_inventory_adjustment_workbook, build_shipment_workbook
from app.inventory_writeback import build_inventory_writeback
from app.destination_country import country_from_text, infer_destination_country


class ParserTests(unittest.TestCase):
    def test_shipment_export_toy_categories_follow_each_item(self):
        task = {"customer": "ZURU", "soNumber": "CAT-TEST", "items": [
            {"product_code": "A", "category": "电子"},
            {"product_code": "B", "category": "塑胶"},
            {"product_code": "C"},
        ]}
        for variant in ("shipping", "warehouse"):
            task["exportVariant"] = variant
            workbook, _ = build_shipment_workbook(task)
            sheet = openpyxl.load_workbook(BytesIO(workbook)).active
            self.assertEqual("玩具类别", sheet["H5"].value)
            self.assertEqual("电子", sheet["H6"].value)
            self.assertEqual("塑胶", sheet["H7"].value)
            self.assertIsNone(sheet["H8"].value)


    def test_destination_country_uses_explicit_country_and_curated_ports(self):
        self.assertEqual("美国", country_from_text("WM US ELWOOD; SAVANNAH"))
        self.assertEqual("英国", country_from_text("YTN-FELIXSTOWE, UK"))
        self.assertEqual("德国", country_from_text("Destination: Hamburg"))
        self.assertEqual("日本", country_from_text("Port of discharge: Yokohama"))
        self.assertEqual("韩国", country_from_text("Destination BUSAN"))

    def test_destination_country_ignores_unrelated_attachment_country(self):
        attachments = [
            {"fields": {"country": "美国"}},
            {"fields": {"destination_port": "FELIXSTOWE, United Kingdom", "country": "英国"}},
        ]
        self.assertEqual("英国", infer_destination_country("走货通知", attachments))
        self.assertEqual("", infer_destination_country("走货通知", [{"fields": {"country": "德国"}}]))

    def test_shipment_export_variants_and_inventory_adjustment(self):
        task = {
            "id": 12, "customer": "ZURU", "soNumber": "45001",
            "plannedShipDate": "2026-08-29", "completedDate": "2026-08-29",
            "containerType": "40HQ", "status": "Completed",
            "items": [{"contract_number": "45001", "product_code": "77772-S001",
                       "product_name": "唱片机 6个/箱", "quantity": 600, "pieces": 100,
                       "warehouse_location": "B栋3楼B1区", "warehouse_match_status": "已匹配",
                       "inventory_customer_name": "THE WAREHOUSE LIMITED",
                       "inventory_country": "新西兰", "inventory_receipt_number": "NO:A2400177"}],
        }
        shipping, _ = build_shipment_workbook(task)
        shipping_sheet = openpyxl.load_workbook(BytesIO(shipping)).active
        self.assertEqual(17, shipping_sheet.max_column)
        self.assertNotIn("放货区", [cell.value for cell in shipping_sheet[5]])

        task["exportVariant"] = "warehouse"
        warehouse, _ = build_shipment_workbook(task)
        warehouse_sheet = openpyxl.load_workbook(BytesIO(warehouse)).active
        self.assertEqual("放货区", warehouse_sheet.cell(5, 18).value)
        self.assertEqual("B栋3楼B1区", warehouse_sheet.cell(6, 18).value)

        adjustment, _ = build_inventory_adjustment_workbook({
            "from": "2026-08-01", "to": "2026-08-31", "tasks": [task]
        })
        adjustment_sheet = openpyxl.load_workbook(BytesIO(adjustment)).active
        self.assertEqual(
            ["洋行", "合同号", "客户名称", "走货国家", "货号", "入库单号", "放货区", "出库日期", "出库数量", "出柜车次"],
            [cell.value for cell in adjustment_sheet[1]],
        )
        self.assertEqual("THE WAREHOUSE LIMITED", adjustment_sheet["C2"].value)
        self.assertEqual("NO:A2400177", adjustment_sheet["F2"].value)
        self.assertEqual("2026-08-29", adjustment_sheet["H2"].value)
        self.assertEqual(600, adjustment_sheet["I2"].value)
        self.assertEqual("待补充", adjustment_sheet["J2"].value)
        self.assertEqual("00FFF2CC", adjustment_sheet["H2"].fill.fgColor.rgb)

        summary, _ = build_completed_shipment_summary_workbook({
            "from": "2026-08-01", "to": "2026-08-31", "tasks": [task]
        })
        summary_sheet = openpyxl.load_workbook(BytesIO(summary), data_only=False).active
        self.assertEqual("走柜任务汇总", summary_sheet.title)
        self.assertEqual("完成日期", summary_sheet["A1"].value)
        self.assertEqual("2026-08-29", summary_sheet["A2"].value)
        self.assertEqual(600, summary_sheet["H2"].value)
        self.assertEqual(100, summary_sheet["I2"].value)

    def test_shipment_export_uses_product_code_factory_and_brand_rules(self):
        task = {
            "customer": "ZURU", "soNumber": "RULES", "items": [
                {"product_code": "92107Q7", "pieces": 10},
                {"product_code": "15756-S001", "pieces": 20},
                {"product_code": "15754UQ1", "pieces": 30},
                {"product_code": "9565GQ1", "pieces": 40},
                {"product_code": "88888", "supplier": "博锐", "pieces": 50},
            ],
        }
        content, _ = build_shipment_workbook(task)
        sheet = openpyxl.load_workbook(BytesIO(content)).active
        self.assertEqual(["华登", "华登", "华康", None, "博锐"], [sheet.cell(row, 1).value for row in range(6, 11)])
        brand_note = sheet.cell(11, 6).value
        self.assertIn("92107品牌：ZURU/BABYCORNS", brand_note)
        self.assertIn("15756品牌：SPONGEBOB SQUAREPANTS", brand_note)
        self.assertIn("157开头的品牌：ZURU/FUGGLER", brand_note)
        self.assertIn("95开头品牌：ZURU/PetsAlive", brand_note)
        self.assertNotIn("88888", brand_note)

    def test_full_container_loading_factory_uses_factory_mix(self):
        for items, expected in (
            ([{"product_code": "92107"}, {"product_code": "9565", "supplier": "兴信"}], "兴信做柜"),
            ([{"product_code": "92107"}, {"product_code": "9298"}], "华登做柜"),
        ):
            content, _ = build_shipment_workbook({"customer": "ZURU", "containerType": "1*40HQ", "items": items})
            sheet = openpyxl.load_workbook(BytesIO(content)).active
            self.assertEqual(expected, sheet["F1"].value)

    def test_shipment_export_leaves_brand_note_empty_when_no_rule_matches(self):
        content, _ = build_shipment_workbook({
            "customer": "ZURU", "soNumber": "NO-BRAND",
            "items": [{"product_code": "88888", "pieces": 1}],
        })
        sheet = openpyxl.load_workbook(BytesIO(content)).active
        self.assertIsNone(sheet.cell(7, 6).value)

    def test_shipment_export_uses_aggregated_order_total_pieces(self):
        content, _ = build_shipment_workbook({
            "customer": "ZURU", "soNumber": "TOTAL",
            "items": [{"product_code": "88888", "pieces": 500, "order_total_pieces": 650}],
        })
        sheet = openpyxl.load_workbook(BytesIO(content)).active
        self.assertEqual(500, sheet.cell(6, 10).value)
        self.assertEqual(650, sheet.cell(6, 15).value)

    def test_shipment_export_fills_every_country_from_task_destination(self):
        content, _ = build_shipment_workbook({
            "customer": "ZURU", "soNumber": "COUNTRY", "destinationCountry": "美国",
            "items": [
                {"product_code": "10001", "country": "法国", "pieces": 1},
                {"product_code": "10002", "pieces": 2},
            ],
        })
        sheet = openpyxl.load_workbook(BytesIO(content)).active
        self.assertEqual(["美国", "美国"], [sheet.cell(row, 7).value for row in range(6, 8)])

    def test_shipment_export_formats_date_control_values(self):
        content, _ = build_shipment_workbook({
            "customer": "ZURU", "soNumber": "DATES", "cutoffDate": "2026-09-15",
            "siDeadline": "2026-09-12T15:00", "items": [],
        })
        sheet = openpyxl.load_workbook(BytesIO(content)).active
        self.assertEqual("SI：9月12日 15:00", sheet.cell(9, 1).value)
        self.assertEqual("截数期：9月15日", sheet.cell(10, 1).value)

    def test_business_inventory_spreadsheet(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "客户库存.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "库存"
            ws.append(["货号", "客户PO", "合同号", "订单数量", "库存数量"])
            ws.append(["77876", "19KM6AYO", "4500211856", 2016, 1800])
            wb.save(path)
            result = parse_business_spreadsheet(path, path.name)
            self.assertEqual("inventory", result["kind"])
            self.assertEqual(1800, result["rows"][0]["inventory_quantity"])

    def test_selected_local_inventory_folder_scan_and_match(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "A车间库存表.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "ZURU"
            ws.append(["洋行", "合同号", "货号", "入库单号", "放货区", "库存数量"])
            ws.append(["ZURU", "4500214717", "77858-S001", "NO:A2513270", "B栋3楼C6区", 2067])
            wb.save(path)
            scan = _scan_inventory_paths([path], "我的库存")
            self.assertEqual("我的库存", scan["folder"])
            self.assertEqual(1, scan["successful_files"])
            match = _match_inventory_paths([path], {"customer":"ZURU", "items":[{
                "contract_number":"4500214717", "product_code":"77858-S001"
            }]})
            self.assertEqual("NO:A2513270", match["items"][0]["receipt_number"])
            self.assertEqual("B栋3楼C6区", match["items"][0]["location"])
            ws["E2"] = "B栋4楼C6区"
            wb.save(path)
            refreshed = _match_inventory_paths([path], {"customer":"ZURU", "items":[{
                "contract_number":"4500214717", "product_code":"77858-S001"
            }]})
            self.assertEqual("B栋4楼C6区", refreshed["items"][0]["location"])

    def test_inventory_writeback_uses_fifo_and_updates_existing_rows(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "库存表.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "ZURU"
            ws.append(["合同号", "货号", "入库日期", "入库单号", "入库数量", "放货区", "出库日期", "出库数量", "出柜车次", "库存数量"])
            ws.append(["45001", "A-S001", "2026-08-01", "NO:1", 100, "A区", None, None, None, 100])
            ws.append(["45001", "A-S001", "2026-08-02", "NO:2", 100, "B区", None, 10, None, 90])
            wb.save(path)
            original = path.read_bytes()
            result = build_inventory_writeback([path], {
                "task_id": 7, "outbound_date": "2026-09-09",
                "items": [{"contract_number":"45001", "product_code":"A-S001", "quantity":150, "outbound_trip":"柜A-01"}],
            })
            self.assertEqual(2, result["updated_rows"])
            self.assertEqual(original, base64.b64decode(result["files"][0]["backup_base64"]))
            updated = openpyxl.load_workbook(BytesIO(base64.b64decode(result["files"][0]["content_base64"]))).active
            self.assertEqual(100, updated["H2"].value)
            self.assertEqual(0, updated["J2"].value)
            self.assertEqual(60, updated["H3"].value)
            self.assertEqual(40, updated["J3"].value)
            self.assertEqual("2026-09-09", updated["G2"].value.date().isoformat())
            self.assertEqual("柜A-01", updated["I2"].value)
            self.assertEqual("柜A-01", updated["I3"].value)

    def test_inventory_writeback_rejects_insufficient_stock_before_editing(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "库存表.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["合同号", "货号", "入库日期", "出库日期", "出库数量", "库存数量"])
            ws.append(["45001", "A-S001", "2026-08-01", None, None, 20])
            wb.save(path)
            original = path.read_bytes()
            with self.assertRaisesRegex(ValueError, "库存不足"):
                build_inventory_writeback([path], {
                    "task_id": 8, "outbound_date": "2026-09-09",
                    "items": [{"contract_number":"45001", "product_code":"A-S001", "quantity":21}],
                })
            self.assertEqual(original, path.read_bytes())

    def test_inventory_writeback_uses_inbound_minus_outbound_when_inventory_is_formula(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "A车间库存表.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "ZURU"
            ws.append(["合同号", "货号", "入库日期", "入库数量", "出库日期", "出库数量", "库存数量"])
            ws.append(["4500214490", "71172-S001", "2026-08-10", 2920, None, None, "=D2-F2"])
            wb.save(path)

            parsed = parse_business_spreadsheet(path, path.name)
            self.assertEqual(1, len(parsed["rows"]))
            self.assertIsNone(parsed["rows"][0]["inventory_quantity"])

            result = build_inventory_writeback([path], {
                "task_id": 9, "outbound_date": "2026-09-09",
                "items": [{"contract_number":"4500214490", "product_code":"71172-S001", "quantity":2920}],
            })
            updated = openpyxl.load_workbook(
                BytesIO(base64.b64decode(result["files"][0]["content_base64"])), data_only=False
            ).active
            self.assertEqual(2920, updated["F2"].value)
            self.assertEqual("=D2-F2", updated["G2"].value)

    def test_inventory_writeback_batches_tasks_and_keeps_latest_date(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "库存表.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["合同号", "货号", "入库日期", "出库日期", "出库数量", "库存数量"])
            ws.append(["45001", "A-S001", "2026-08-01", "2026-09-05", 5, 100])
            wb.save(path)
            result = build_inventory_writeback([path], {
                "task_id": "batch", "outbound_date": "2026-09-09",
                "items": [
                    {"contract_number":"45001", "product_code":"A-S001", "quantity":10,
                     "writeback_task_id":"7", "outbound_date":"2026-09-01"},
                    {"contract_number":"45001", "product_code":"A-S001", "quantity":20,
                     "writeback_task_id":"8", "outbound_date":"2026-09-03"},
                ],
            })
            updated = openpyxl.load_workbook(BytesIO(base64.b64decode(result["files"][0]["content_base64"]))).active
            self.assertEqual(35, updated["E2"].value)
            self.assertEqual(70, updated["F2"].value)
            self.assertEqual("2026-09-05", updated["D2"].value.date().isoformat())
            self.assertEqual(["7", "8"], [allocation["writeback_task_id"] for allocation in result["allocations"]])

    def test_inventory_scan_excludes_writeback_backups_and_deduction_reports(self):
        with tempfile.TemporaryDirectory() as root:
            folder = Path(root)
            (folder / "A车间库存表.xlsx").touch()
            (folder / "A车间库存表.回写前备份-20260909.xlsx").touch()
            (folder / "库存扣减写回表_2026-09.xlsx").touch()
            self.assertEqual(["A车间库存表.xlsx"], [path.name for path in _inventory_paths(folder)])

    def test_real_inspection_headers_choose_effective_result(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "兴信每日验货总结.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "8月份"
            ws.append(["兴信每日验货总结表"])
            ws.append(["日期", "验货地点", "客户名称", "第三方", "合同编号", "客户/PO", "货号", "产品名称", "数量", "箱数", "洋行 结果", "第三方结果", "HOLD/REJ 原因"])
            ws.append(["2026-08-01", "兴信", "ZURU", "WM自检", 4500207414, "0997413614", "157149-S001", "产品", 2834, 1417, "NA", "PASS", None])
            wb.save(path)
            result = parse_business_spreadsheet(path, path.name)
            self.assertEqual("inspection", result["kind"])
            self.assertEqual("PASS", result["rows"][0]["inspection_result"])
            self.assertEqual("0997413614", result["rows"][0]["customer_po"])

    def test_body_fields(self):
        result = parse_body("SO# SHZ7952680\n柜型 1*40'HQ\nSI：3月27日 09:00\n盐田港\n装柜完成后必须拉网保护货物")
        self.assertEqual("SHZ7952680", result["so_number"])
        self.assertEqual("1*40HQ", result["container_type"])
        self.assertEqual(f"{datetime.now().year}-03-27T09:00", result["si_deadline"])
        self.assertEqual("盐田", result["port"])
        self.assertEqual("装柜完成后必须拉网保护货物", result["special_requirements"])

    def test_shipping_deadline_formats_and_aliases(self):
        cases = [
            ("截单时间：2026-07-20 09:00", "2026-07-20T09:00"),
            ("最迟截补料：7/20早上8点", "2026-07-20T08:00"),
            ("SI & VGM CUT OFF TIME: 21-JUL 10:00", "2026-07-21T10:00"),
            ("截SI / VGM/ ENS时间：16-Jul-2026 10:00", "2026-07-16T10:00"),
            ("SI cut off date/time: 2026/07/21 15:00:00", "2026-07-21T15:00"),
        ]
        for body, expected in cases:
            with self.subTest(body=body):
                self.assertEqual(expected, parse_body(body)["si_deadline"])

    def test_cutoff_formats_and_aliases(self):
        cases = [
            ("CY CUT: 2026-07-23 12:00", "2026-07-23T12:00"),
            ("CY CLOSING:20-Jul-2026 12:00", "2026-07-20T12:00"),
            ("port cut-off is July 10, 2026", "2026-07-10T00:00"),
            ("截重柜时间：2026/07/22 12:00", "2026-07-22T12:00"),
        ]
        for body, expected in cases:
            with self.subTest(body=body):
                self.assertEqual(expected, parse_body(body)["cutoff_date"])

    def test_earliest_current_deadline_wins_and_history_is_only_fallback(self):
        result = parse_body(
            "CFS closing: 20-Jul 15:00\nCFS closing: 17-Jul 15:00\n"
            "发送时间: 2026年7月1日\nCY CLOSING: 16-Jul 12:00"
        )
        self.assertEqual("2026-07-17T15:00", result["cutoff_date"])
        self.assertEqual("2026-07-16T12:00", result["_fallback_cutoff_date"])
        self.assertIn("识别到多个时间", result["_date_warnings"][0])

    def test_spaced_named_deadline_keeps_day_and_time(self):
        result = parse_body("截补料时间：23- JUL 15:00")
        self.assertEqual("2026-07-23T15:00", result["si_deadline"])

    def test_subject_so_wins_and_booking_header_does_not_capture_next_line(self):
        result = parse_body(
            "S/O  CAN/SBYTS1074680827B - 兴信\n"
            "SO NUMBER\nBooking No.\nDelivery\nDelivery date\nBN260501700"
        )
        self.assertEqual("CAN/SBYTS1074680827B", result["so_number"])

    def test_booking_number_must_be_on_same_line(self):
        result = parse_body("Booking No.\nDelivery\nBN260501700")
        self.assertEqual("BN260501700", result["so_number"])

    def test_body_fields_keep_complete_requirement_lines(self):
        result = parse_body(
            "请务必做好分货措施，要么每票货物的外箱都贴对应的SO号码。\n"
            "装车的时候按每个SO的单号去装货，不要3个SO的货混装，多装、漏装。\n"
            "普通问候语"
        )
        self.assertEqual(
            "请务必做好分货措施，要么每票货物的外箱都贴对应的SO号码。\n"
            "装车的时候按每个SO的单号去装货，不要3个SO的货混装，多装、漏装。",
            result["special_requirements"],
        )

    def test_special_requirements_exclude_admin_history_and_isolated_fallbacks(self):
        result = parse_body(
            "请查收附件SO，开仓后请尽早安排打单提柜，谢谢！\n"
            "截补料：7/20 15:00，请关注船期动态。\n"
            "报关资料请发送到 shipping@example.com\n"
            "拉网\n拍照\n立放\n"
            "From: Old Sender\n装柜完成后必须拉网并拍照"
        )
        self.assertEqual("", result["special_requirements"])

    def test_special_requirements_keep_complete_operational_sentences(self):
        result = parse_body(
            "装柜要求：\n"
            "所有展示架要求立放，装完柜后请拍摄装满柜及半关柜门照片。\n"
            "不同SO不得混装，每票货物之间请用纸皮隔开。\n"
            "电子委托代码4400000000，报关费由工厂承担。"
        )
        self.assertEqual(
            "所有展示架要求立放，装完柜后请拍摄装满柜及半关柜门照片。\n"
            "不同SO不得混装，每票货物之间请用纸皮隔开。",
            result["special_requirements"],
        )

    def test_amazon_requirement_uses_exact_current_paragraph(self):
        paragraph = (
            "亚马逊不接受：铁皮箱，双开门 /one way container；"
            "如有异常请先不要提柜并申请换柜，谢谢！"
        )
        result = parse_body(f"请留意装柜注意事项。\n\n{paragraph}\n\n报关资料请发给货代")
        self.assertEqual(paragraph, result["special_requirements"])

    def test_shipment_export_splits_top_and_bottom_notes(self):
        content, _ = build_shipment_workbook({
            "customer": "ZURU", "soNumber": "NOTES",
            "specialRequirements": "兴信拖柜，博锐送兴信拼柜，深圳报关\n按SO分货，不能混装，箱唛朝向柜门",
            "items": [{"product_code": "88888", "supplier": "兴信"}, {"product_code": "77772", "supplier": "博锐"}],
        })
        sheet = openpyxl.load_workbook(BytesIO(content)).active
        self.assertEqual("兴信拖柜，博锐送兴信拼柜，深圳报关", sheet["F1"].value)
        self.assertEqual("按SO分货，不能混装，箱唛朝向柜门", next(cell.value for cell in sheet["N"] if cell.value == "按SO分货，不能混装，箱唛朝向柜门"))

    def test_external_factory_does_not_export_bottom_email_note(self):
        content, _ = build_shipment_workbook({
            "customer": "ZURU", "soNumber": "EXTERNAL",
            "specialRequirements": "按SO分货，不能混装",
            "items": [{"product_code": "88888", "supplier": "博锐"}],
        })
        sheet = openpyxl.load_workbook(BytesIO(content)).active
        self.assertEqual("送博锐拼柜，深圳报关", sheet["F1"].value)
        self.assertNotIn("按SO分货", [cell.value for row in sheet.iter_rows() for cell in row])

    def test_customer_pickup_note_is_normalized(self):
        content, _ = build_shipment_workbook({
            "customer": "ZURU", "soNumber": "PICKUP",
            "specialRequirements": "宏达报关行，客上车\n按SO分货",
            "items": [{"product_code": "88888", "supplier": "兴信"}],
        })
        sheet = openpyxl.load_workbook(BytesIO(content)).active
        self.assertEqual("宏达报关行，客上柜，拼柜，深圳报关", sheet["F1"].value)
        self.assertEqual("按SO分货", next(cell.value for cell in sheet["N"] if cell.value == "按SO分货"))

    def test_eml_and_attachment_isolation(self):
        msg = EmailMessage()
        msg["Subject"] = "订舱通知"
        msg["From"] = "customer@example.com"
        msg.set_content("Booking No: HKGG40351600")
        msg.add_attachment(b"book", maintype="application", subtype="pdf", filename="booking.pdf")
        with tempfile.TemporaryDirectory() as root:
            result = parse_eml(msg.as_bytes(), Path(root) / "one")
            self.assertEqual("订舱通知", result["subject"])
            self.assertEqual(1, len(result["attachments"]))
            self.assertTrue((Path(root) / "one" / result["attachments"][0]["stored_name"]).exists())

    def test_email_skips_one_unreadable_attachment_and_keeps_valid_packing_list(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Packing list"
        sheet.append(["SKU NO.", "Shipping Quantity", "NO of CARTON"])
        sheet.append(["15754", 4068, 1017])
        valid_attachment = BytesIO()
        workbook.save(valid_attachment)

        message = EmailMessage()
        message["Subject"] = "S/O 兴信"
        message["From"] = "customer@example.com"
        message.set_content("Booking notice")
        message.add_attachment(valid_attachment.getvalue(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename="packing list.xlsx")
        message.add_attachment(b"encrypted workbook", maintype="application", subtype="vnd.ms-excel", filename="交仓指引.xls")

        upload = UploadFile(file=BytesIO(message.as_bytes()), filename="2.eml")
        result = asyncio.run(parse_email_batch([upload]))

        self.assertEqual(1, result["parsed"])
        self.assertEqual(0, result["failed"])
        self.assertEqual("15754", result["items"][0]["items"][0]["product_code"])
        self.assertIn("已跳过无法解析附件：交仓指引.xls", result["items"][0]["warnings"][0])
        self.assertEqual("failed", result["items"][0]["attachment_results"][1]["status"])

    def test_standard_packing_list(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "PL.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Packing list"
            ws.append(["SKU NO.", "Shipping Quantity", "NO of CARTON", "CBM", "Customer PO#", "ZURU PO No.", "L", "W", "H"])
            ws.append(["15754", 4068, 1017, 15.5, "0958543467", "4500189918", 44, 36.5, 27.5])
            wb.save(path)
            result = parse_attachment(path, path.name)
            self.assertEqual("packing_list", result["kind"])
            self.assertEqual("15754", result["items"][0]["product_code"])
            self.assertEqual(1017, result["items"][0]["pieces"])
            self.assertEqual("44*36.5*27.5", result["items"][0]["box_dimensions"])

    def test_amazon_allocation_parses_and_splits_loading_factories(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "EditLineItems.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Edit Line Items-Skyler"
            ws.append(["PO", "Spec#", "Title", "Expected Quantity", "Case Pack", "Total Carton",
                       "Supplier", "45#", "total CBM", "New Booking", "New Booking Key#", "Note"])
            ws.append(["A1", "111-S001", "产品1", 120, 12, 10, "Danli", "45001", 1.2,
                       "Booking 6-1*40HQ-1", "AMZ992N219710SZ1", None])
            ws.append(["A2", "222-S001", "产品2", 240, 12, 20, "Royal", "45002", 2.4,
                       "Booking 6-1*40HQ-2", "AMZ992N219710SZ1", "旧资料写华康做柜"])
            ws.append(["A3", "333-S001", "产品3", 360, 12, 30, "Royal", "45003", 3.6,
                       "Booking 6-1*45HQ-1", "AMZ992N219710SZ1", "旧资料写华康做柜"])
            wb.save(path)

            parsed = parse_attachment(path, path.name)
            self.assertEqual("amazon_allocation", parsed["kind"])
            self.assertEqual(3, len(parsed["items"]))
            groups = build_amazon_shipment_groups(
                "Booking# AMZ992N219710SZ1",
                "1*40HQ 丹尼做柜\n1*40HQ 兴信做柜\n1*45HQ 兴信做柜",
                [{"kind": parsed["kind"], "items": parsed["items"]}],
            )
            self.assertEqual(
                [("1*40HQ", "丹尼"), ("1*40HQ", "兴信"), ("1*45HQ", "兴信")],
                [(group["container_type"], group["loading_factory"]) for group in groups],
            )

    def test_retail_unit_is_quantity_and_cartons_are_pieces(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "英国分柜箱单.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Packing list"
            ws.append(["SKU NO.", "DESCRIPTION ", "Shipping QUANTITY", "Retail Unit", "PC PER\nCARTON", "NO of\nCARTON"])
            ws.append(["MB208", "MIXED IN CTN,2PCS/CTN", 1636, 3272, 2, 1636])
            wb.save(path)
            item = parse_attachment(path, path.name)["items"][0]
            self.assertEqual(2, item["spec"])
            self.assertEqual(3272, item["quantity"])
            self.assertEqual(1636, item["pieces"])

    def test_product_name_and_carton_spec_are_separate(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "中文货名.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Packing list"
            ws.append(["货号", "货名", "数量", "件数"])
            ws.append(["92105", "尿布蛋12个/箱", 2400, 200])
            ws.append(["A100", "魔法小鸟", 500, 50])
            wb.save(path)
            items = parse_attachment(path, path.name)["items"]
            self.assertEqual("尿布蛋", items[0]["product_name"])
            self.assertEqual(12, items[0]["spec"])
            self.assertEqual(200, items[0]["pieces"])
            self.assertEqual("魔法小鸟", items[1]["product_name"])
            self.assertIsNone(items[1]["spec"])
            self.assertEqual(50, items[1]["pieces"])

    def test_pieces_are_calculated_only_when_carton_column_is_missing(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "无箱数.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Packing list"
            ws.append(["货号", "货名", "数量", "每箱"])
            ws.append(["A100", "大脑", 2000, 20])
            wb.save(path)
            item = parse_attachment(path, path.name)["items"][0]
            self.assertEqual(100, item["pieces"])

    def test_loading_factory_marker_keeps_all_factories(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "multi-factory.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Packing list"
            ws.append(["SKU NO.", "Actual Factory", "Quantity", "Remarks"])
            ws.append(["A1", "兴信", 10, ""])
            ws.append(["B1", "泰亨", 20, "40HQ 兴信装柜"])
            ws.append(["汇总", "", 30, ""])
            wb.save(path)
            result = parse_attachment(path, path.name)
            filtered = filter_items_for_email(
                result["items"], "standard", result["fields"].get("loading_factory", "")
            )
            self.assertEqual("兴信", result["fields"]["loading_factory"])
            self.assertEqual(["A1", "B1"], [item["product_code"] for item in filtered])

    def test_multi_factory_without_loading_marker_keeps_only_xingxin(self):
        items = [{"product_code": "A1", "supplier": "兴信"},
                 {"product_code": "B1", "supplier": "泰亨"}]
        filtered = filter_items_for_email(items, "standard")
        self.assertEqual(["A1"], [item["product_code"] for item in filtered])

    def test_full_container_keeps_huadeng_cargo(self):
        items = [{"product_code": "92107", "supplier": "华登"},
                 {"product_code": "9565", "supplier": "兴信"}]
        self.assertEqual(items, filter_items_for_email(items, "standard"))
        self.assertEqual(items[:1], filter_items_for_email(items[:1], "standard"))

    def test_warehouse_delivery_keeps_external_factory_cargo(self):
        items = [{"product_code": "A1", "supplier": "兴信"},
                 {"product_code": "B1", "supplier": "泰亨"}]
        filtered = filter_items_for_email(items, "warehouse")
        self.assertEqual(["A1", "B1"], [item["product_code"] for item in filtered])

    def test_other_factories_are_not_kept_as_xingxin_work(self):
        items = [{"product_code": "A1", "supplier": "丹尼", "factory_remark": "华康"},
                 {"product_code": "B1", "supplier": "泰亨", "factory_remark": "库有"}]
        self.assertEqual([], filter_items_for_email(items, "standard"))

    def test_description_keeps_specific_product_name_and_removes_packing_text(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "产品名称.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Packing list"
            ws.append(["SKU NO.", "DESCRIPTION", "Retail Unit", "PC PER CARTON", "NO of CARTON"])
            ws.append(["15792SLD1", "S001-FUGGLER-BIG NUGGETS-75PCS/CTN", 3150, 75, 42])
            ws.append(["MB208", "ROBO ALIVE-DINOSAUR AND PETS ALIVE-MAGIC BIRD-2PCS/CTN", 3272, 2, 1636])
            wb.save(path)
            items = parse_attachment(path, path.name)["items"]
            self.assertEqual("非凡系列 BIG NUGGETS", items[0]["product_name"])
            self.assertEqual("恐龙 / 魔法小鸟", items[1]["product_name"])

    def test_product_names_are_extracted_from_packing_list_description(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "Packing List.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Packing list"
            ws.append(["SKU NO.", "DESCRIPTION", "Retail Unit", "PC PER CARTON", "NO of CARTON"])
            ws.append(["MEC416", "S001-EUR-MTS-SLIME MART-MIXED-SERIES1-SHOPPING BAG AND SHOPPING BASKET M IXED IN ENDCAP,32PCS/ENDCAP,1ENDCAP/CTN", 2016, 1, 63])
            ws.append(["MSQ52", "S001-MTS-MIX BRAND-TOYS-MIXED-SERIES 1-5 SURPRISE-A-LOT-A AXOLOTLS/PUPPY PET SHOP/FUGGLER-BABY PINKLES KEYRINGS/RAINBOCORNS-EGGZANIA/OOSH-SLIME/SNACKLES-KEYCHAINS MIXED IN PDQ,41PCS/PDQ/CTN", 4920, 1, 120])
            ws.append(["77772GQ1", "S002-MTS-MINI BRANDS-VINYL-SERIES 1-CAPSULE,13PCS/PDQ/CTN", 13000, 13, 1000])
            ws.append(["77843GQ1", "S002-MTS-MINI BRANDS-KAWAII KITCHEN-SERIES 1-CAPSULE,13PCS/PDQ/CTN", 4810, 13, 370])
            ws.append(["18101UQ1", "S002-MTS-NOOK NOOKS-NOOK NOOKS-SERIES 1-BLIND BOX-COLOR BOX,6PCS/PDQ/CTN", 6000, 6, 1000])
            ws.append(["92120TQ1", "S002-NB-MTS-RAINBOCORNS-SPARKLE HEART SURPRISE-SERIES 1-AXOLOTLCORN SURPRISE-SHRINK WRAP,6PCS/PDQ/CTN", 4638, 6, 773])
            wb.save(path)
            items = parse_attachment(path, path.name)["items"]
            self.assertEqual("SLIME MART SHOPPING BAG AND SHOPPING BASKET", items[0]["product_name"])
            self.assertEqual(32, items[0]["spec"])
            self.assertIn("非凡系列 BABY PINKLES", items[1]["product_name"])
            self.assertEqual(41, items[1]["spec"])
            self.assertEqual("MINI BRANDS VINYL", items[2]["product_name"])
            self.assertEqual("MINI BRANDS KAWAII KITCHEN", items[3]["product_name"])
            self.assertEqual("NOOK NOOKS", items[4]["product_name"])
            self.assertEqual("RAINBOCORNS SPARKLE HEART SURPRISE AXOLOTLCORN SURPRISE", items[5]["product_name"])

    def test_yax_notice(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "YAX5634105.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["马士基FCA并柜拖车通知单"])
            ws.append(["船公司订舱单号", "149601102372"])
            ws.append(["柜型", "40GP(40HQ SUB)"])
            ws.append(["截关时间", "2026/3/7 12:00"])
            ws.append(["PO号码", "工厂", "SO", "CBM"])
            ws.append(["10001519913", "YAT", "YAX5634105", 13.935])
            wb.save(path)
            result = parse_attachment(path, path.name)
            self.assertEqual("yax", result["kind"])
            self.assertEqual("149601102372", result["fields"]["booking_number"])
            self.assertEqual("1*40GP", result["fields"]["container_type"])
            self.assertEqual("YAX5634105", result["items"][0]["so_number"])

    def test_word_warehouse(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "warehouse.docx"
            document = Document()
            document.add_paragraph("仓库地址：深圳市盐田区测试仓")
            document.add_paragraph("SO# SHZ123456")
            document.save(path)
            result = parse_attachment(path, path.name)
            self.assertEqual("深圳市盐田区测试仓", result["fields"]["warehouse"])
            self.assertEqual("SHZ123456", result["fields"]["so_number"])

    def test_multi_warehouse_grouping(self):
        parsed = [
            {"filename": "1000001-SO.pdf", "fields": {"delivery_address": "ELWOOD"}, "items": []},
            {"filename": "1000001 PKL.xlsx", "fields": {}, "items": [{"product_code": "A"}]},
            {"filename": "1000002-SO.pdf", "fields": {"delivery_address": "ELWOOD"}, "items": []},
            {"filename": "1000002 PKL.xlsx", "fields": {}, "items": [{"product_code": "B"}]},
            {"filename": "1000003-SO.pdf", "fields": {"delivery_address": "SAVANNAH"}, "items": []},
            {"filename": "1000003 PKL.xlsx", "fields": {}, "items": [{"product_code": "C"}]},
        ]
        groups = build_warehouse_groups(parsed)
        self.assertEqual(2, len(groups))
        elwood = next(group for group in groups if group["warehouse"] == "ELWOOD")
        self.assertEqual(["1000001", "1000002"], elwood["references"])
        self.assertEqual(2, len(elwood["items"]))

    def test_header_priority_and_total_filter_support(self):
        columns = _column_map(["ITEM NO.", "SKU NO.", "NO CARTON"])
        self.assertEqual(1, columns["product_code"])
        self.assertEqual(2, columns["pieces"])

    def test_rejects_pdf_label_as_address(self):
        self.assertEqual("", _clean_address("Movement Term: Port - Port"))
        self.assertEqual("ELWOOD STORAGE", _clean_address("ELWOOD STORAGE"))


if __name__ == "__main__":
    unittest.main()
