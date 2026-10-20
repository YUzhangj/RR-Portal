import unittest
from io import BytesIO
from openpyxl import load_workbook
from app.shipment_export import build_shipment_workbook

class IncompleteExportTests(unittest.TestCase):
 def test_missing_unit_and_total_are_marked_for_both_factories(self):
  for company in ('Xingxin','Huadeng'):
   task={'company':company,'customer':'ZURU','exportMissingWeights':[{'itemIndex':0,'field':'net_weight_per_box','totalField':'net_weight'}],'items':[{'product_code':'A','quantity':12,'pieces':1,'gross_weight_per_box':2}]}
   content,name=build_shipment_workbook(task);s=load_workbook(BytesIO(content)).active
   self.assertEqual('待核对',s['Q6'].value);self.assertEqual('待核对',s['L6'].value);self.assertIn('待核对',name)
   totals=[row for row in s.iter_rows() if any(c.value=='合计' for c in row)]
   self.assertEqual('待核对',s.cell(totals[0][0].row,12).value)
   self.assertTrue(str(s['K6'].value).startswith('='))
 def test_huadeng_pallet_rows_and_customer_item_columns_remain_aligned(self):
  task={'company':'Huadeng','customer':'ZURU','exportTemplate':{'company':'Huadeng','layoutKey':'CustomerItem','weightRule':'Exact'},'exportDetails':{'shipmentMode':'Warehouse'},'exportMissingWeights':[{'itemIndex':1,'field':'net_weight_per_box','totalField':'net_weight'}],'items':[{'product_code':'A','pieces':1,'pallet_count':1,'gross_weight_per_box':2,'net_weight_per_box':1},{'product_code':'B','pieces':2,'gross_weight_per_box':3}]}
  content,_=build_shipment_workbook(task);s=load_workbook(BytesIO(content)).active
  self.assertEqual('待核对',s['R8'].value);self.assertEqual('待核对',s['L8'].value)
  self.assertEqual(1,s['K7'].value);self.assertEqual(1,s['R6'].value)
 def test_complete_export_keeps_formulas(self):
  task={'company':'Xingxin','customer':'ZURU','items':[{'product_code':'A','pieces':2,'gross_weight_per_box':3,'net_weight_per_box':1}]}
  content,name=build_shipment_workbook(task);s=load_workbook(BytesIO(content)).active
  self.assertEqual('=INT(J6*Q6)',s['L6'].value);self.assertNotIn('待核对',name)
