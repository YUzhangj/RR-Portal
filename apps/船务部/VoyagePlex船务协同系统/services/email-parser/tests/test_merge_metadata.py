"""Bounded worksheet-merge XML and Huadeng business-equivalence fixtures."""
from io import BytesIO
from pathlib import Path
import tempfile
import tracemalloc
from types import SimpleNamespace
import unittest
from unittest.mock import patch
import zipfile
import openpyxl

from app import attachments as m
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ERROR='XLSX merge metadata is invalid or exceeds safe limits'

class Archive:
    def __init__(self,xml,size=None,compressed=None):
        self.xml=xml;self.size=len(xml) if size is None else size
        self.compressed=self.size if compressed is None else compressed
        self.reads=[]
    def getinfo(self,path):return SimpleNamespace(file_size=self.size,compress_size=self.compressed)
    def open(self,path):
        stream=BytesIO(self.xml);read=stream.read
        def bounded(n):self.reads.append(n);return read(n)
        stream.read=bounded;return stream

def ranges(xml,**kwargs):
    archive=Archive(xml,**kwargs)
    return m._sheet_merge_ranges(SimpleNamespace(_archive=archive),SimpleNamespace(_worksheet_path='xl/worksheets/sheet1.xml'))

def xml(refs,prefix=''):
    tag=(prefix+':') if prefix else ''
    xmlns='xmlns:'+prefix if prefix else 'xmlns'
    return ('<'+tag+'worksheet '+xmlns+'="'+NS+'"><'+tag+'sheetData/>'
        '<'+tag+'mergeCells>'+''.join('<'+tag+'mergeCell ref="'+r+'"/>' for r in refs)
        +'</'+tag+'mergeCells></'+tag+'worksheet>').encode()

class MergeXML(unittest.TestCase):
    def test_namespaces_multiple_columns_absolute_refs_and_dedup(self):
        refs=['D2:D3','E2:F3','J2:J3','N2:P3','D2:D3','$Q$4:$S$5','A1']
        for prefix in ['','s','sheet']:
            with self.subTest(prefix=prefix):
                actual=ranges(xml(refs,prefix));expected={(1,3,3,4),(1,3,4,6),(1,3,9,10),(1,3,13,16),(3,5,16,19),(0,1,0,1)}
                self.assertEqual(set(actual),expected);self.assertEqual(len(actual),len(expected))
        self.assertEqual(set(ranges(xml(refs).decode().encode('utf-16'))),expected)
    def test_full_openpyxl_set_order_and_ranges_equal_on_multiple_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'synthetic.xlsx';wb=openpyxl.Workbook()
            wb.active.title='Other';sheet=wb.create_sheet('分柜明细')
            for ref in ['D2:D3','E2:F3','J2:J3','N2:P3','Q4:S5','A1']:sheet.merge_cells(ref)
            wb.save(path);wb.close()
            full=openpyxl.load_workbook(path,read_only=False);stream=openpyxl.load_workbook(path,read_only=True)
            try:
                expected=[(r.min_row-1,r.max_row,r.min_col-1,r.max_col) for r in full['分柜明细'].merged_cells.ranges]
                self.assertEqual(m._sheet_merge_ranges(stream,stream['分柜明细']),expected)
                self.assertEqual(m._sheet_merge_ranges(stream,stream['Other']),[])
            finally:full.close();stream.close()
    def test_malformed_references_xml_structure_and_dtd_fail_static(self):
        bad=[xml(['A0:B2']),xml(['A1:XFE2']),xml(['A1:B1048577']),xml(['B2:A1']),xml(['Sheet!A1:B2']),
            b'<worksheet xmlns="'+NS.encode()+b'"><mergeCells>',
            xml(['A1:B2']).replace(b'<mergeCell ',b'<foreign:mergeCell xmlns:foreign="urn:wrong" '),
            b'<!DOCTYPE worksheet [<!ENTITY x "private-data">]><worksheet xmlns="'+NS.encode()+b'">&x;</worksheet>',
            b'<worksheet xmlns="urn:wrong"/>',
            b'<worksheet xmlns="'+NS.encode()+b'"><mergeCell ref="A1:B2"/></worksheet>']
        for value in bad:
            with self.subTest(length=len(value)),self.assertRaisesRegex(ValueError,'^'+ERROR+'$'):ranges(value)
    def test_depth_merge_count_zip_metadata_and_unknown_entry_bounds(self):
        root=b'<worksheet xmlns="'+NS.encode()+b'">'
        for value,kwargs in [(root+b'<x>'*64+b'</x>'*64+b'</worksheet>',{}),
                (xml(['A1:B2']*10001),{}),(xml([]),{'size':64*1024*1024+1}),
                (xml([]),{'compressed':64*1024*1024+1}),(xml([]),{'size':0})]:
            with self.subTest(length=len(value),kwargs=kwargs),self.assertRaisesRegex(ValueError,'^'+ERROR+'$'):ranges(value,**kwargs)
    def test_huge_formatted_empty_rows_do_not_accumulate_a_tree(self):
        value=b'<worksheet xmlns="'+NS.encode()+b'"><sheetData>'+b'<row><c s="1"/></row>'*150000+b'</sheetData><mergeCells><mergeCell ref="D2:F3"/></mergeCells></worksheet>'
        archive=Archive(value);tracemalloc.start()
        try:
            result=m._sheet_merge_ranges(SimpleNamespace(_archive=archive),SimpleNamespace(_worksheet_path='sheet.xml'))
            _,peak=tracemalloc.get_traced_memory()
        finally:tracemalloc.stop()
        self.assertEqual(result,[(1,3,3,6)]);self.assertLessEqual(max(archive.reads),32768)
        self.assertLess(peak,2*1024*1024) # XML input was allocated before tracing.
    def test_memory_error_is_never_hidden_as_an_empty_merge_list(self):
        archive=Archive(xml([]))
        with patch.object(archive,'open',side_effect=MemoryError()):
            with self.assertRaises(MemoryError):m._sheet_merge_ranges(SimpleNamespace(_archive=archive),SimpleNamespace(_worksheet_path='sheet.xml'))
    def test_actual_decompressed_reads_and_node_count_are_bounded(self):
        class EndlessWhitespace:
            def __init__(self):self.first=True;self.total=0;self.max_read=0
            def __enter__(self):return self
            def __exit__(self,*_):pass
            def read(self,n):
                self.max_read=max(self.max_read,n)
                value=b'<worksheet xmlns="'+NS.encode()+b'">' if self.first else b' '*n
                self.first=False;self.total+=len(value);return value
        archive=Archive(xml([]));stream=EndlessWhitespace()
        with patch.object(archive,'open',return_value=stream),self.assertRaisesRegex(ValueError,'^'+ERROR+'$'):
            m._sheet_merge_ranges(SimpleNamespace(_archive=archive),SimpleNamespace(_worksheet_path='sheet.xml'))
        self.assertEqual(stream.total,64*1024*1024+1);self.assertLessEqual(stream.max_read,32768)
        value=b'<worksheet xmlns="'+NS.encode()+b'">'+b'<x/>'*1000000+b'</worksheet>'
        with self.assertRaisesRegex(ValueError,'^'+ERROR+'$'):ranges(value)
        with self.assertRaisesRegex(ValueError,'^'+ERROR+'$'):ranges(xml([]),size=len(xml([]))+1)

class HuadengBusiness(unittest.TestCase):
    def test_mixed_cartons_all_business_fields_equal_full_book_metadata(self):
        from app.huadeng_forms import parse_huadeng_form
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'synthetic.xlsx';wb=openpyxl.Workbook();wb.active.title='Other'
            ws=wb.create_sheet('华登分柜明细')
            rows=[['备注','序号','洋行','合同','货号','货名','国家','玩具类别','数量','件数','毛重','净重','体积','客人PO#','每单总件数','每箱毛重','每箱净重'],
                ['华登',1,'ZURU','4500208073','MSLD200','第一款','美国','塑胶',2058,42,2772,1337.7,60.627,'PO',84,66,31.85],
                ['华登',2,'ZURU',None,None,'第二款','美国','塑胶',882,None,None,652.68,None,None,None,None,15.54]]
            for row in rows:ws.append(row)
            for ref in ['D2:D3','E2:E3','J2:J3','K2:K3','M2:M3','N2:N3','O2:O3','P2:P3']:ws.merge_cells(ref)
            wb.save(path);wb.close()
            full=openpyxl.load_workbook(path,data_only=True,read_only=False)
            read=openpyxl.load_workbook(path,data_only=True,read_only=True)
            try:
                values=m._read_sheet_rows(read['华登分柜明细'])
                old=parse_huadeng_form(values,[(r.min_row-1,r.max_row,r.min_col-1,r.max_col) for r in full['华登分柜明细'].merged_cells.ranges])
            finally:full.close();read.close()
            loader=m.load_values_workbook;calls=[]
            def only_streaming(*args,**kwargs):
                calls.append(kwargs);book=loader(*args,**kwargs);self.assertTrue(book.read_only);return book
            with patch.object(m,'load_values_workbook',side_effect=only_streaming):actual=m.parse_excel(path)
            self.assertEqual(actual,old);self.assertEqual(len(calls),1)
            self.assertEqual(len(actual['items']),2);self.assertEqual(actual['items'][1]['product_code'],'MSLD200')
            self.assertIsNone(actual['items'][1]['pieces']);self.assertEqual(actual['items'][0]['packing_group'],actual['items'][1]['packing_group'])

if __name__=='__main__':unittest.main(verbosity=2)
