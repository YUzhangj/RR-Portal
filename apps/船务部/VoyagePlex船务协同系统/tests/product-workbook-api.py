"""End-to-end workbook preview/commit regression with isolated API, parser and SQLite."""
from io import BytesIO
from pathlib import Path
import http.cookiejar,json,os,secrets,socket,sqlite3,subprocess,sys,tempfile,time,urllib.error,urllib.request
from openpyxl import Workbook
ROOT=Path(__file__).resolve().parents[1]
def port():
 with socket.socket() as s:s.bind(('127.0.0.1',0));return s.getsockname()[1]
def run():
 with tempfile.TemporaryDirectory(prefix='voyageplex-product-import-') as temp:
  temp=Path(temp);db=temp/'test.db';api_port,parser_port=port(),port();base=f'http://127.0.0.1:{api_port}';processes=[];logs=[]
  jar=http.cookiejar.CookieJar();opener=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
  env={k:v for k,v in os.environ.items() if not k.startswith(('VOYAGEPLEX_MAIL_','VOYAGEPLEX_HUADENG_MAIL_'))}
  def start(command,cwd,values,name):
   log=open(temp/f'{name}.log','wb');logs.append(log);proc=subprocess.Popen(command,cwd=cwd,env=dict(env,**values),stdout=log,stderr=log);processes.append(proc);return proc
  def wait(url):
   for _ in range(200):
    if any(p.poll() is not None for p in processes):raise AssertionError('Test service stopped')
    try:
     with urllib.request.urlopen(url,timeout=1):return
    except OSError:time.sleep(.1)
   raise AssertionError('Test service startup timeout')
  def req(path,body=None,status=200):
   data=None if body is None else json.dumps(body).encode();request=urllib.request.Request(base+path,data=data,headers={'Content-Type':'application/json'} if data else {})
   try:r=opener.open(request,timeout=20)
   except urllib.error.HTTPError as e:r=e
   value=json.loads(r.read());assert r.code==status,(path,r.code,value);return value
  def preview(data):
   boundary='boundary-'+secrets.token_hex(10)
   raw=(f'--{boundary}\r\nContent-Disposition: form-data; name="files"; filename="huadeng.xlsx"\r\nContent-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n'.encode()+data+f'\r\n--{boundary}--\r\n'.encode())
   request=urllib.request.Request(base+'/api/product-infos/workbooks/preview',data=raw,headers={'Content-Type':f'multipart/form-data; boundary={boundary}'})
   with opener.open(request,timeout=20) as response:return json.load(response)['rows']
  try:
   python=sys.executable
   start([str(python),'-m','uvicorn','app.main:app','--host','127.0.0.1','--port',str(parser_port)],ROOT/'services/email-parser',{},'parser')
   wait(f'http://127.0.0.1:{parser_port}/health')
   start(['dotnet',str(ROOT/'server/VoyagePlex.Api/bin/Debug/net8.0/VoyagePlex.Api.dll')],ROOT/'server/VoyagePlex.Api',{'ConnectionStrings__Default':f'Data Source={db}','ASPNETCORE_URLS':base,'EmailParser__BaseUrl':f'http://127.0.0.1:{parser_port}'},'api')
   wait(base+'/api/auth/setup-status');password=secrets.token_urlsafe(16)
   req('/api/auth/setup',{'username':'admin','displayName':'Test admin','password':password},201);req('/api/auth/login',{'username':'admin','password':password})
   book=Workbook();sheet=book.active;sheet.title='621柜'
   headers=['备注','序号','洋行','合同','货号','货名','国家','玩具类别','数量','件数','毛重','净重','体积','客PO','每单总件数','每箱毛重','每箱净重']
   a=['华登',1,'ZURU','C1','IMPORT-HUA-A','猫蛋6个/箱','美国','塑胶',60,10,999,888,1.5,'PO',10,2.65,1.09]
   sheet.append(headers);sheet.append(a)
   other=book.create_sheet('313车');other.append(headers[:13]+['箱唛CM','','','客PO','每单总数量','每箱毛重','每箱净净重']);other.append(['C',1,'PWP','C2','IMPORT-HUA-B','小老虎400/箱','美国','搪胶',3200,8,999,888,.28,33.02,27.305,40.005,'PO',8,13.33,11.3])
   duplicate=book.create_sheet('重复产品');duplicate.append(headers);duplicate.append(a)
   stream=BytesIO();book.save(stream);data=stream.getvalue();rows=preview(data)
   assert len(rows)==3 and {row['sheetName'] for row in rows}=={'621柜','313车','重复产品'}
   assert any('重复' in warning for row in rows for warning in row['warnings'])
   with sqlite3.connect(db) as c:assert c.execute("SELECT COUNT(*) FROM ProductInfos WHERE ProductCode LIKE 'IMPORT-HUA-%'").fetchone()[0]==0
   chosen=[row for row in rows if row['sheetName']!='重复产品']
   assert req('/api/product-infos/workbooks/commit',{'rows':chosen})=={'added':2,'updated':0,'skipped':0}
   again=preview(data);chosen=[row for row in again if row['sheetName']!='重复产品'];assert all(row['existingId'] for row in chosen)
   assert req('/api/product-infos/workbooks/commit',{'rows':chosen})=={'added':0,'updated':0,'skipped':2}
   selected=next(row for row in chosen if row['productCode']=='IMPORT-HUA-A');selected.update(updateExisting=True,grossWeightPerBox=4)
   assert req('/api/product-infos/workbooks/commit',{'rows':[selected]})['updated']==1
   with sqlite3.connect(db) as c:
    a=c.execute("SELECT GrossWeightPerBox,FactoryRemark FROM ProductInfos WHERE ProductCode='IMPORT-HUA-A'").fetchone();b=c.execute("SELECT NetWeightPerBox,NetNetWeightPerBox,BoxDimensions FROM ProductInfos WHERE ProductCode='IMPORT-HUA-B'").fetchone()
    assert float(a[0])==4 and a[1]=='';assert b[0] is None and float(b[1])==11.3 and b[2]=='33.02*27.305*40.005'
   print('Workbook API passed: real parser, multi-sheet preview, no write before confirmation, duplicate warnings, net-net weights/dimensions and explicit existing-product update.')
  finally:
   for p in reversed(processes):p.terminate();p.wait(timeout=15)
   for log in logs:log.close()
if __name__=='__main__':run()
